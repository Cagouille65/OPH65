import io
import html
import os
import json
import ssl
import smtplib
import urllib.parse
import base64
import mimetypes
import threading
import time
import zipfile
import tempfile
import shutil
import gzip
import secrets
from copy import copy
from datetime import date, datetime

import pandas as pd
import streamlit as st
import plotly.express as px
import streamlit.components.v1 as components
from email.message import EmailMessage
from openpyxl import load_workbook

try:
    import qrcode
    HAS_QRCODE = True
except Exception:
    qrcode = None
    HAS_QRCODE = False

try:
    from supabase import create_client
    HAS_SUPABASE = True
except Exception:
    create_client = None
    HAS_SUPABASE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, KeepTogether
    HAS_REPORTLAB = True
except Exception:
    HAS_REPORTLAB = False

try:
    import keyring
    HAS_KEYRING = True
except Exception:
    keyring = None
    HAS_KEYRING = False

try:
    from st_aggrid import AgGrid, GridOptionsBuilder, JsCode
    HAS_AGGRID = True
except Exception:
    HAS_AGGRID = False

APP_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(APP_DIR, "logo.jpg")
CONFIG_PATH = os.path.join(APP_DIR, "oph65_config.json")
REFERENCE_DB_PATH = os.path.join(APP_DIR, "oph65_logements.json")
KEYRING_SERVICE = "OPH65_Suivi_Travaux_SMTP"

st.set_page_config(page_title="OPH65 | Suivi des travaux", page_icon="🏗️", layout="wide")

# Mode Cloud : activé automatiquement sur Streamlit Community Cloud, ou via [app].cloud_mode dans Secrets.
def _secret(path, default=None):
    try:
        cur = st.secrets
        for part in path.split("."):
            cur = cur[part]
        return cur
    except Exception:
        return default

CLOUD_MODE = bool(_secret("app.cloud_mode", False)) or bool(os.environ.get("STREAMLIT_SHARING_MODE"))


@st.cache_resource
def get_supabase_client():
    """Client Supabase côté serveur. La clé n'est jamais exposée au navigateur."""
    if not HAS_SUPABASE:
        return None
    url = _secret("supabase.url", "")
    key = _secret("supabase.secret_key", "") or _secret("supabase.service_role_key", "") or _secret("supabase.key", "")
    if not url or not key:
        return None
    try:
        return create_client(str(url), str(key))
    except Exception:
        return None


def supabase_ready():
    return get_supabase_client() is not None


def _db_json_value(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except Exception:
        pass
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(v).isoformat()
    if hasattr(v, "item"):
        try: return v.item()
        except Exception: pass
    return v


def _serialize_row(row):
    return {str(k): _db_json_value(v) for k, v in dict(row).items()}


def db_save_settings(cfg):
    client = get_supabase_client()
    if not client: return
    try:
        client.table("app_settings").upsert({"key":"config", "value":cfg, "updated_at":datetime.utcnow().isoformat()}).execute()
    except Exception:
        pass


def db_load_settings():
    client = get_supabase_client()
    if not client: return {}
    try:
        res = client.table("app_settings").select("value").eq("key","config").limit(1).execute()
        rows = getattr(res, "data", None) or []
        return rows[0].get("value", {}) if rows else {}
    except Exception:
        return {}


def db_upsert_tracking(df):
    client = get_supabase_client()
    if not client or df is None or df.empty or "N° LOT" not in df.columns: return
    records=[]
    now=datetime.utcnow().isoformat()
    for _, row in df.iterrows():
        lot=normalize_lot_key(row.get("N° LOT"))
        if lot:
            records.append({"lot_no":lot,"data":_serialize_row(row),"updated_at":now})
    for i in range(0,len(records),400):
        client.table("suivi_logements").upsert(records[i:i+400]).execute()


def db_load_tracking():
    client=get_supabase_client()
    if not client: return pd.DataFrame()
    all_rows=[]
    start=0
    try:
        while True:
            res=client.table("suivi_logements").select("data").range(start,start+999).execute()
            rows=getattr(res,"data",None) or []
            all_rows.extend([r.get("data",{}) for r in rows if isinstance(r.get("data"),dict)])
            if len(rows)<1000: break
            start+=1000
    except Exception:
        return pd.DataFrame()
    if not all_rows: return pd.DataFrame()
    out=pd.DataFrame(all_rows)
    for c in out.columns:
        if _looks_like_date_column(c): out[c]=pd.to_datetime(out[c],errors="coerce")
    return calculate_planned_dates(apply_reference_database(out))


def db_save_reference_database(database):
    client=get_supabase_client()
    if not client: return
    now=datetime.utcnow().isoformat(); records=[]
    for lot,row in (database or {}).items():
        key=normalize_lot_key(lot)
        if key: records.append({"lot_no":key,"data":_reference_row(row,key),"updated_at":now})
    for i in range(0,len(records),400):
        client.table("logements").upsert(records[i:i+400]).execute()


def db_load_reference_database():
    client=get_supabase_client()
    if not client: return {}
    data={}; start=0
    try:
        while True:
            res=client.table("logements").select("lot_no,data").range(start,start+999).execute()
            rows=getattr(res,"data",None) or []
            for r in rows:
                key=normalize_lot_key(r.get("lot_no")); row=r.get("data") or {}
                if key: data[key]=_reference_row(row,key)
            if len(rows)<1000: break
            start+=1000
    except Exception:
        return {}
    return data


def db_save_excel_bytes(content, filename):
    client=get_supabase_client()
    if not client or not content: return
    packed=base64.b64encode(gzip.compress(content, compresslevel=6)).decode("ascii")
    client.table("app_files").upsert({"key":"source_workbook","filename":filename or "Suivi_Travaux_OPH65.xlsm","content_b64":packed,"updated_at":datetime.utcnow().isoformat()}).execute()


def db_load_excel_bytes():
    client=get_supabase_client()
    if not client: return None, ""
    try:
        res=client.table("app_files").select("filename,content_b64").eq("key","source_workbook").limit(1).execute()
        rows=getattr(res,"data",None) or []
        if not rows: return None,""
        raw=gzip.decompress(base64.b64decode(rows[0]["content_b64"]))
        return raw, rows[0].get("filename") or "Suivi_Travaux_OPH65.xlsm"
    except Exception:
        return None,""


def db_log_mail(kind, recipient, subject, lot_no="", bc_no="", work=""):
    client=get_supabase_client()
    if not client:return
    try:
        client.table("mail_history").insert({"kind":kind,"recipient":recipient,"subject":subject,"lot_no":str(lot_no or ""),"bc_no":str(bc_no or ""),"work_category":str(work or ""),"sent_at":datetime.utcnow().isoformat()}).execute()
    except Exception:
        pass


def _shutdown_application(delay=2.0):
    """Arrête Streamlit après avoir figé l'interface sur un écran de fermeture."""
    time.sleep(delay)
    os._exit(0)


@st.dialog("Quitter l’application ?")
def confirm_application_exit():
    st.write(
        "Cette action va arrêter complètement OPH65 et fermer la fenêtre "
        "d’invite de commandes ouverte au démarrage."
    )
    st.caption("Pensez à enregistrer vos éventuelles modifications avant de quitter.")
    c1, c2 = st.columns(2)
    if c1.button("Annuler", use_container_width=True):
        st.rerun()
    if c2.button("🚪 Quitter maintenant", type="primary", use_container_width=True):
        # Quitter d'abord la page Streamlit évite l'écran « Connection error »
        # affiché par le front-end lorsque le serveur disparaît brutalement.
        threading.Thread(target=_shutdown_application, kwargs={"delay": 2.0}, daemon=True).start()
        components.html(
            """
            <script>
            try {
                const doc = window.parent.document;
                doc.body.innerHTML = `
                  <div style="font-family:Arial,sans-serif;display:flex;align-items:center;justify-content:center;height:100vh;background:#f5f7fb;">
                    <div style="text-align:center;background:white;padding:42px 56px;border-radius:16px;box-shadow:0 8px 28px rgba(0,0,0,.12)">
                      <div style="font-size:42px">✅</div>
                      <h2 style="color:#174f91;margin:12px 0">OPH65 est fermé</h2>
                      <p style="color:#4b5563">Vous pouvez fermer cet onglet.</p>
                    </div>
                  </div>`;
            } catch (e) {}
            </script>
            """,
            height=0,
            width=0,
        )
        st.stop()

WORK_CATEGORIES = [
    "NETTOYAGE AV TRVX", "POLYVALENT", "CONFORMITE ELECTRICITE/GAZ", "SOL",
    "PLOMBERIE", "MENUISERIE", "PEINTURE", "NETTOYAGE FIN DE CHANTIER",
]
WORK_DISPLAY_NAMES = {
    "NETTOYAGE AV TRVX": "Nettoyage avant travaux",
    "POLYVALENT": "Polyvalent",
    "CONFORMITE ELECTRICITE/GAZ": "Élec / Gaz",
    "SOL": "Sol",
    "PLOMBERIE": "Plomberie",
    "MENUISERIE": "Menuiserie",
    "PEINTURE": "Peinture",
    "NETTOYAGE FIN DE CHANTIER": "Nettoyage fin de chantier",
}
DISPLAY_TO_CATEGORY = {v: k for k, v in WORK_DISPLAY_NAMES.items()}
REFERENCE_COLUMNS = ["N° LOT", "GDC", "RSD", "Bât", "Ent", "Porte"]
BASE_EDITABLE = ["N° LOT", "GDC", "DATE SAISIE ELS", "PAL"]
STATUS_LIST = ["Terminé", "En cours", "En retard", "Non démarré"]
DEFAULT_STATUS_COLORS = {
    "Terminé": "#C6EFCE",
    "En cours": "#BDD7EE",
    "En retard": "#FFC7CE",
    "Non démarré": "#D9D9D9",
}
DEFAULT_PRIMARY_COLOR = "#174f91"
DEFAULT_SIDEBAR_COLOR = "#123e73"
DATE_FORMAT_OPTIONS = {
    "JJ/MM/AAAA (31/12/2025)": "DD/MM/YYYY",
    "AAAA-MM-JJ (2025-12-31)": "YYYY-MM-DD",
    "JJ-MM-AAAA (31-12-2025)": "DD-MM-YYYY",
    "MM/JJ/AAAA (12/31/2025)": "MM/DD/YYYY",
}


def work_cols(category):
    return {
        "order": f"{category} - N° COMMANDE",
        "start": f"{category} - DATE SAISIE CMD",
        "planned": f"{category} - PRÉVI. J+5 ouvrés",
        "real": f"{category} - DATE RÉELLE FIN TRAVAUX",
    }

WORK_COLUMNS = {WORK_DISPLAY_NAMES[c]: work_cols(c) for c in WORK_CATEGORIES}

DISPLAY_HEADERS = {
    "DATE SAISIE ELS": "DATE ELS",
    "Ent": "Entrée",
}
_DISPLAY_WORK_HEADERS = {
    # Les retours à la ligne sont volontaires : les en-têtes métier doivent tenir sur 2 lignes.
    "NETTOYAGE AV TRVX": ("NET AV TRX\nN°BC", "DATE BC NETTOYAGE\nAV TRX", "DATE THEO NET\nAV TRX", "DATE REELLE NET\nA V TRX"),
    "POLYVALENT": ("POLYVALENT\nN°BC", "DATE BC\nPOLYVALENT", "DATE THEO\nPOLYVALENT", "DATE REELLE\nPOLYVALENT"),
    "CONFORMITE ELECTRICITE/GAZ": ("CONFORMITE ELEC/GAZ\nN°BC", "DATE BC CONFORMITE\nELEC/GAZ", "DATE THEO CONFORMITE\nELEC/GAZ", "DATE REELLE CONFORMITE\nELEC/GAZ"),
    "SOL": ("SOL\nN°BC", "DATE BC\nSOL", "DATE THEO\nSOL", "DATE REELLE\nSOL"),
    "PLOMBERIE": ("PLOMBERIE\nN°BC", "DATE BC\nPLOMBERIE", "DATE THEO\nPLOMBERIE", "DATE REELLE\nPLOMBERIE"),
    "MENUISERIE": ("MENUISERIE\nN°BC", "DATE BC\nMENUISERIE", "DATE THEO\nMENUISERIE", "DATE REELLE\nMENUISERIE"),
    "PEINTURE": ("PEINTURE\nN°BC", "DATE BC\nPEINTURE", "DATE THEO\nPEINTURE", "DATE REELLE\nPEINTURE"),
    "NETTOYAGE FIN DE CHANTIER": ("NETTOYAGE\nN°BC", "DATE BC\nNETTOYAGE", "DATE THEO\nNETTOYAGE", "DATE REELLE\nNETTOYAGE"),
}
for _cat, _labels in _DISPLAY_WORK_HEADERS.items():
    _cols = work_cols(_cat)
    for _key, _label in zip(("order", "start", "planned", "real"), _labels):
        DISPLAY_HEADERS[_cols[_key]] = _label

def display_header(name):
    return DISPLAY_HEADERS.get(name, name)


def load_config():
    defaults = {
        "source_path": "",
        "auto_save": True,
        "residences": [],
        "lot_emails": {},
        "primary_color": DEFAULT_PRIMARY_COLOR,
        "sidebar_color": DEFAULT_SIDEBAR_COLOR,
        "status_colors": DEFAULT_STATUS_COLORS.copy(),
        "date_format": "JJ/MM/AAAA (31/12/2025)",
        "username": "Administrateur",
        "users": ["Administrateur"],
        "smtp_server": "smtp.gmail.com",
        "smtp_port": 465,
        "smtp_use_ssl": True,
        "smtp_auth_required": True,
        "smtp_use_starttls": False,
        "smtp_email": "",
        "email_recipients": "",
        "analysis_email_subject_template": "OPH65 – Synthèse d'avancement des travaux – {date}",
        "analysis_email_body": "Bonjour,\n\nVeuillez trouver ci-joint la synthèse d'avancement des travaux de l'ensemble des logements suivis.\n\nCordialement,\n\nM. Cédric SCHMIT\nOffice Public de l'Habitat des Hautes-Pyrénées\nResponsable Technique de Secteur TN2",
        "delay_email_subject_template": "Retard travaux – lot {lot} – BC n° {bc}",
        "delay_email_body_template": "Bonjour,\n\nSauf erreur de notre part, l'échéance contractuelle des travaux correspondant au bon de commande n° {bc}, relatif au lot « {lot} », est arrivée à son terme. À ce jour, nous constatons que les prestations commandées ne sont pas achevées.\n\nJe vous remercie de bien vouloir me confirmer, par retour de courriel, si les travaux sont désormais terminés. Dans le cas contraire, merci de m'indiquer la nouvelle date ferme d'intervention et d'achèvement prévue.\n\nJe vous rappelle que les délais d'exécution sont fixés conformément aux pièces contractuelles, notamment au CCTP accepté lors de la commande. Tout dépassement non justifié est susceptible d'entraîner l'application des pénalités de retard prévues au marché, sous réserve des stipulations contractuelles applicables.\n\nCordialement,\n\nM. Cédric SCHMIT\nOffice Public de l'Habitat des Hautes-Pyrénées\nResponsable Technique de Secteur TN2",
        "background_image": "",
        "background_overlay": 0.22,
        "background_mode": "Couleur",
        "background_color": "#f5f7fb",
        "font_family": "Arial",
        "font_color": "#1f2937",
    }
    try:
        if os.path.exists(CONFIG_PATH):
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                saved = json.load(f)
            if isinstance(saved, dict):
                defaults.update(saved)
    except Exception:
        pass

    # Sur le Cloud, les réglages persistants Supabase complètent le fichier local.
    if CLOUD_MODE and supabase_ready():
        try:
            saved_cloud = db_load_settings()
            if isinstance(saved_cloud, dict): defaults.update(saved_cloud)
        except Exception:
            pass

    # Les secrets Streamlit ont priorité sur les réglages locaux pour les paramètres sensibles.
    secret_map = {
        "smtp_server": "smtp.server",
        "smtp_port": "smtp.port",
        "smtp_use_ssl": "smtp.use_ssl",
        "smtp_auth_required": "smtp.auth_required",
        "smtp_use_starttls": "smtp.use_starttls",
        "smtp_email": "smtp.email",
        "email_recipients": "smtp.default_recipients",
    }
    for cfg_key, secret_key in secret_map.items():
        value = _secret(secret_key, None)
        if value is not None:
            defaults[cfg_key] = value
    return defaults


def save_config(**updates):
    cfg = load_config()
    cfg.update(updates)
    # Ne jamais écrire le mot de passe SMTP sur disque.
    cfg.pop("smtp_password", None)
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    if CLOUD_MODE and supabase_ready():
        db_save_settings(cfg)


def normalize_lot_key(value):
    """Normalise le N° LOT utilisé comme clé unique du référentiel logements."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def _reference_row(row, key):
    """Normalise une fiche logement du référentiel local."""
    row = row if isinstance(row, dict) else {}
    return {
        "N° LOT": key,
        "RSD": _clean(row.get("RSD", "")),
        "Bât": _clean(row.get("Bât", "")),
        "Ent": _clean(row.get("Ent", row.get("Entrée", ""))),
        "Porte": _clean(row.get("Porte", "")),
        "Résidence": _clean(row.get("Résidence", "")),
        "Adresse": _clean(row.get("Adresse", "")),
        "Suite adresse": _clean(row.get("Suite adresse", row.get("Adresse Suite", ""))),
        "Code postal": _clean(row.get("Code postal", row.get("CPostal", ""))),
        "Ville": _clean(row.get("Ville", row.get("Localité", ""))),
        "Typologie": _clean(row.get("Typologie", "")),
        "Secteur": _clean(row.get("Secteur", "")),
        "Date mise en service": _clean(row.get("Date mise en service", row.get("Date Mise En Service", ""))),
        "Type chauffage": _clean(row.get("Type chauffage", row.get("Mode de Chauffage", ""))),
        "Contrat PES": _clean(row.get("Contrat PES", "")),
    }


def load_reference_database():
    """Charge le référentiel maître. Supabase est prioritaire en mode Cloud."""
    if CLOUD_MODE and supabase_ready():
        cloud = db_load_reference_database()
        if cloud:
            return cloud
    try:
        if os.path.exists(REFERENCE_DB_PATH):
            with open(REFERENCE_DB_PATH, "r", encoding="utf-8") as f:
                raw = json.load(f)
            if isinstance(raw, dict):
                cleaned = {}
                for lot, row in raw.items():
                    key = normalize_lot_key(lot)
                    if key and isinstance(row, dict):
                        cleaned[key] = _reference_row(row, key)
                # Initialise automatiquement Supabase au premier démarrage.
                if CLOUD_MODE and supabase_ready() and cleaned:
                    try: db_save_reference_database(cleaned)
                    except Exception: pass
                return cleaned
    except Exception:
        pass
    return {}


def save_reference_database(database):
    """Enregistre atomiquement le référentiel logements dans le dossier de l'application."""
    payload = {}
    for lot, row in (database or {}).items():
        key = normalize_lot_key(lot)
        if key:
            payload[key] = _reference_row(row, key)
    tmp = REFERENCE_DB_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2, default=str)
    os.replace(tmp, REFERENCE_DB_PATH)
    if CLOUD_MODE and supabase_ready():
        db_save_reference_database(payload)


def lookup_lot_reference(lot_no):
    key = normalize_lot_key(lot_no)
    if not key:
        return None
    return load_reference_database().get(key)


def apply_reference_database(df):
    """Applique le référentiel maître à un DataFrame en utilisant N° LOT comme clé."""
    if df is None or df.empty or "N° LOT" not in df.columns:
        return df.copy() if isinstance(df, pd.DataFrame) else df
    database = load_reference_database()
    if not database:
        return df.copy()
    out = df.copy()
    keys = out["N° LOT"].map(normalize_lot_key)
    for col in ["RSD", "Bât", "Ent", "Porte"]:
        mapped = keys.map(lambda k: database.get(k, {}).get(col, None))
        if col not in out.columns:
            out[col] = mapped
        else:
            has_ref = keys.map(lambda k: k in database)
            out.loc[has_ref, col] = mapped.loc[has_ref]
    return out


def reference_database_from_excel(file):
    """Extrait les fiches d'identité depuis l'onglet List LGMT d'un fichier Excel/XLSM."""
    lgmt = pd.read_excel(file, sheet_name="List LGMT")
    if lgmt.empty:
        raise ValueError("L'onglet 'List LGMT' est vide.")
    # Accepte les variantes de libellés présentes dans les fichiers OPH65.
    def val(row, *names):
        for name in names:
            if name in lgmt.columns:
                return row.get(name, "")
        return ""
    database = {}
    for _, row in lgmt.iterrows():
        key = normalize_lot_key(val(row, "N° Lot", "N° LOT", "N° lot"))
        if not key:
            continue
        date_service = val(row, "Date Mise En Service", "Date mise en service")
        try:
            if pd.notna(date_service) and date_service != "":
                date_service = pd.to_datetime(date_service).strftime("%Y-%m-%d")
            else:
                date_service = ""
        except Exception:
            date_service = _clean(date_service)
        database[key] = _reference_row({
            "RSD": val(row, "Rsd", "RSD"),
            "Bât": val(row, "Bât.", "Bât"),
            "Ent": val(row, "Ent.", "Ent", "Entrée"),
            "Porte": val(row, "N° Porte", "Porte"),
            "Résidence": val(row, "Résidence"),
            "Adresse": val(row, "Rue", "Adresse"),
            "Suite adresse": val(row, "Adresse Suite", "Suite adresse"),
            "Code postal": val(row, "CPostal", "Code postal"),
            "Ville": val(row, "Localité", "Ville"),
            "Typologie": val(row, "Typologie"),
            "Secteur": val(row, "Secteur"),
            "Date mise en service": date_service,
            "Type chauffage": val(row, "Mode de Chauffage", "Type chauffage"),
            "Contrat PES": val(row, "Contrat PES"),
        }, key)
    return database

def load_secure_smtp_password(sender_email):
    """Relit le secret SMTP depuis Streamlit Secrets (Cloud), puis le coffre local si disponible."""
    cloud_password = _secret("smtp.password", "")
    if cloud_password:
        return str(cloud_password)
    if not sender_email or not HAS_KEYRING:
        return ""
    try:
        return keyring.get_password(KEYRING_SERVICE, sender_email) or ""
    except Exception:
        return ""


def save_secure_smtp_password(sender_email, password):
    """Mémorise le mot de passe SMTP localement ; en Cloud, utiliser Streamlit Secrets."""
    if CLOUD_MODE:
        return bool(_secret("smtp.password", "")) or not bool(password)
    if not sender_email or not HAS_KEYRING:
        return False
    try:
        if password:
            keyring.set_password(KEYRING_SERVICE, sender_email, password)
        else:
            try:
                keyring.delete_password(KEYRING_SERVICE, sender_email)
            except Exception:
                pass
        return True
    except Exception:
        return False


def save_background_file(uploaded_file):
    """Enregistre l'image de fond dans le dossier de l'application."""
    if uploaded_file is None:
        return ""
    ext = os.path.splitext(uploaded_file.name or "")[1].lower()
    if ext not in (".jpg", ".jpeg", ".png", ".webp"):
        ext = ".jpg"
    filename = f"oph65_background{ext}"
    for old_ext in (".jpg", ".jpeg", ".png", ".webp"):
        old = os.path.join(APP_DIR, f"oph65_background{old_ext}")
        if os.path.exists(old) and os.path.basename(old) != filename:
            try:
                os.remove(old)
            except Exception:
                pass
    with open(os.path.join(APP_DIR, filename), "wb") as f:
        f.write(uploaded_file.getvalue())
    return filename


def delete_background_file(filename):
    if not filename:
        return
    path = os.path.join(APP_DIR, os.path.basename(filename))
    try:
        if os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def background_css(filename, overlay=0.22, mode="Couleur", color="#f5f7fb"):
    """CSS plein écran responsive : l'image couvre l'écran sans déformation."""
    if mode == "Couleur" or not filename:
        safe_color = color if isinstance(color, str) and color.startswith("#") else "#f5f7fb"
        return f".stApp {{ background:{safe_color}; }}"
    path = os.path.join(APP_DIR, os.path.basename(filename))
    if not os.path.exists(path):
        safe_color = color if isinstance(color, str) and color.startswith("#") else "#f5f7fb"
        return f".stApp {{ background:{safe_color}; }}"
    try:
        with open(path, "rb") as f:
            encoded = base64.b64encode(f.read()).decode("ascii")
        mime = mimetypes.guess_type(path)[0] or "image/jpeg"
        alpha = min(max(float(overlay), 0.0), 0.95)
        return f"""
        .stApp {{
            background-image: linear-gradient(rgba(245,247,251,{alpha}), rgba(245,247,251,{alpha})),
                              url('data:{mime};base64,{encoded}');
            background-size: cover;
            background-position: center center;
            background-repeat: no-repeat;
            background-attachment: fixed;
            min-height: 100vh;
        }}
        [data-testid='stMainBlockContainer'] {{
            background: rgba(255,255,255,0.08);
            border-radius: 14px;
        }}
        """
    except Exception:
        return ".stApp { background:#f5f7fb; }"


def _clean(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).replace("\n", " ").strip()


APARTMENT_REF_SUBHEADERS = {"RSD", "Bât", "Ent", "Porte"}
KNOWN_SUBHEADERS = APARTMENT_REF_SUBHEADERS | {
    "N° COMMANDE", "DATE SAISIE CMD", "PRÉVI. J+5 ouvrés", "DATE RÉELLE FIN TRAVAUX"
}


def normalize_columns(raw_row1, raw_row2):
    row1 = [_clean(v) for v in raw_row1]
    row2 = [_clean(v) for v in raw_row2] if raw_row2 is not None else []
    two_row_header = sum(1 for v in row2 if v in KNOWN_SUBHEADERS) >= 4
    headers, seen = [], {}
    if two_row_header:
        categories, last = [], ""
        for v in row1:
            if v:
                last = v
            categories.append(last)
        for i, (cat, sub) in enumerate(zip(categories, row2)):
            # Colonnes fixes du classeur OPH65 : noms stables, même si H1 contient une formule dynamique.
            if i == 0: name = "N° LOT"
            elif i == 1: name = "GDC"
            elif i == 2: name = "RSD"
            elif i == 3: name = "Bât"
            elif i == 4: name = "Ent"
            elif i == 5: name = "Porte"
            elif i == 6: name = "DATE SAISIE ELS"
            elif i == 7: name = "PAL"
            elif i == 40: name = "NDJI (J ELS→PAL)"
            elif sub:
                name = f"{cat} - {sub}"
            else:
                name = cat or "Colonne"
            seen[name] = seen.get(name, 0) + 1
            headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    else:
        for v in row1:
            name = v or "Colonne"
            seen[name] = seen.get(name, 0) + 1
            headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return headers, (2 if two_row_header else 1)


def _looks_like_date_column(name):
    n = name.upper()
    if "N° COMMANDE" in n or n in ("N° LOT", "GDC", "NDJI (J ELS→PAL)"):
        return False
    return "DATE" in n or "PRÉVI" in n or n == "PAL"


def add_business_days(value, days=5):
    if value is None or pd.isna(value):
        return pd.NaT
    return (pd.Timestamp(value).normalize() + pd.offsets.BDay(days)).to_pydatetime()


def calculate_planned_dates(df):
    out = df.copy()
    for work, cols in WORK_COLUMNS.items():
        if cols["start"] in out.columns:
            start = pd.to_datetime(out[cols["start"]], errors="coerce")
            computed = start.map(lambda x: add_business_days(x, 5) if pd.notna(x) else pd.NaT)
            if cols["planned"] not in out.columns:
                out[cols["planned"]] = computed
            else:
                # La date prévisionnelle est une donnée calculée : on la recalcule systématiquement.
                out[cols["planned"]] = computed
    return out


def _fill_apartment_references_from_list_lgmt(df, excel_file):
    """Renseigne RSD/Bât/Entrée/Porte depuis `List LGMT` par N° LOT.

    Dans le classeur OPH65 ces champs sont souvent des formules VLOOKUP. Pandas ne
    recalcule pas Excel ; on reproduit donc la jointure pour garantir leur présence
    dans l'interface, les synthèses et les PDF.
    """
    if df.empty or "N° LOT" not in df.columns or "List LGMT" not in excel_file.sheet_names:
        return df
    try:
        lgmt = pd.read_excel(excel_file, sheet_name="List LGMT", header=None)
        if lgmt.empty or lgmt.shape[1] < 12:
            return df

        # Correspondance identique aux VLOOKUP du classeur : A→clé, D→RSD, E→Bât, F→Entrée, L→Porte.
        ref = lgmt.iloc[:, [0, 3, 4, 5, 11]].copy()
        ref.columns = ["N° LOT", "RSD", "Bât", "Ent", "Porte"]

        ref["_lot_key"] = ref["N° LOT"].map(normalize_lot_key)
        ref = ref[ref["_lot_key"] != ""].drop_duplicates("_lot_key", keep="first")
        lookup = ref.set_index("_lot_key")[["RSD", "Bât", "Ent", "Porte"]]
        keys = df["N° LOT"].map(normalize_lot_key)

        out = df.copy()
        for col in ["RSD", "Bât", "Ent", "Porte"]:
            mapped = keys.map(lookup[col])
            if col not in out.columns:
                out[col] = mapped
            else:
                current = out[col]
                missing = current.isna() | current.astype(str).str.strip().isin(["", "nan", "None"])
                out.loc[missing, col] = mapped.loc[missing]
        return out
    except Exception:
        # Le chargement principal doit rester utilisable même avec un onglet List LGMT atypique.
        return df


def read_excel(file):
    # ExcelFile permet de lire également List LGMT sans dépendre des valeurs mises en cache des formules.
    excel_file = pd.ExcelFile(file)
    raw = pd.read_excel(excel_file, header=None, sheet_name=0)
    if raw.empty:
        return pd.DataFrame()
    headers, start = normalize_columns(raw.iloc[0].tolist(), raw.iloc[1].tolist() if len(raw) > 1 else None)
    df = raw.iloc[start:].copy()
    df.columns = headers[:len(df.columns)]
    df = df.dropna(how="all").reset_index(drop=True)
    df = _fill_apartment_references_from_list_lgmt(df, excel_file)
    df = apply_reference_database(df)
    for c in df.columns:
        if _looks_like_date_column(c):
            df[c] = pd.to_datetime(df[c], errors="coerce")
    return calculate_planned_dates(df)


def status_for(row, work):
    cols = WORK_COLUMNS[work]
    real = row.get(cols["real"], pd.NaT)
    if pd.notna(real):
        return "Terminé"
    planned = row.get(cols["planned"], pd.NaT)
    if pd.notna(planned):
        return "En retard" if pd.Timestamp(planned).date() < date.today() else "En cours"
    start = row.get(cols["start"], pd.NaT)
    order = row.get(cols["order"], None)
    if pd.notna(start) or (pd.notna(order) and str(order).strip()):
        return "En cours"
    return "Non démarré"


def enrich(df):
    if df.empty:
        return df.copy()
    out = apply_reference_database(calculate_planned_dates(df))
    for work in WORK_COLUMNS:
        out[f"Statut {work}"] = out.apply(lambda r: status_for(r, work), axis=1)
    status_cols = [f"Statut {w}" for w in WORK_COLUMNS]

    def global_status(r):
        vals = [r.get(c, "Non démarré") for c in status_cols]
        if "En retard" in vals:
            return "En retard"
        if all(v == "Terminé" for v in vals):
            return "Terminé"
        if any(v in ("En cours", "Terminé") for v in vals):
            return "En cours"
        return "Non démarré"

    out["Statut global"] = out.apply(global_status, axis=1)
    return out


def get_date_columns(df):
    return [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]


def get_column_config(df):
    fmt = DATE_FORMAT_OPTIONS.get(st.session_state.date_format, "DD/MM/YYYY")
    return {c: st.column_config.DateColumn(display_header(c).replace("\n", " "), format=fmt) for c in get_date_columns(df)}


def coerce_value_for_column(df, column, value):
    """Convertit une valeur de formulaire vers le dtype existant sans provoquer d'upcast Pandas."""
    if column not in df.columns:
        return value
    if _looks_like_date_column(column) or pd.api.types.is_datetime64_any_dtype(df[column]):
        if value in (None, "") or pd.isna(value):
            return pd.NaT
        return pd.Timestamp(value)
    return value


def normalize_date_dtypes(df):
    out = df.copy()
    for c in out.columns:
        if _looks_like_date_column(c):
            out[c] = pd.to_datetime(out[c], errors="coerce")
    return out


def clean_scalar(v):
    if v is None or pd.isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v.item() if hasattr(v, "item") else v


def save_to_source_xlsm(df, path):
    """Met à jour le classeur d'origine de façon atomique, sans supprimer VBA, feuilles, styles ni formules."""
    if not path or not os.path.exists(path):
        raise FileNotFoundError("Le fichier Excel d'origine est introuvable. Rechargez-le depuis Import / Export.")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm") and not zipfile.is_zipfile(path):
        raise ValueError(
            "Le fichier configuré n'est plus un classeur Excel OOXML valide (XLSX/XLSM). "
            "Il est probablement endommagé ou ne correspond pas au fichier actuellement chargé. "
            "Rechargez le classeur original depuis Import / Export avant d'enregistrer."
        )
    keep_vba = ext in (".xlsm", ".xltm")
    wb = load_workbook(path, keep_vba=keep_vba, data_only=False)
    ws = wb[wb.sheetnames[0]]
    col_map = {
        "N° LOT": 1, "GDC": 2, "RSD": 3, "Bât": 4, "Ent": 5, "Porte": 6,
        "DATE SAISIE ELS": 7, "PAL": 8, "NDJI (J ELS→PAL)": 41,
    }
    for idx, cat in enumerate(WORK_CATEGORIES):
        base = 9 + idx * 4
        cols = work_cols(cat)
        col_map[cols["order"]] = base
        col_map[cols["start"]] = base + 1
        col_map[cols["planned"]] = base + 2
        col_map[cols["real"]] = base + 3

    template_row = 3
    for i, (_, row) in enumerate(df.reset_index(drop=True).iterrows(), start=3):
        if i > ws.max_row:
            ws.insert_rows(i)
            for col in range(1, 42):
                src, dst = ws.cell(template_row, col), ws.cell(i, col)
                if src.has_style: dst._style = copy(src._style)
                if src.number_format: dst.number_format = src.number_format
                if src.alignment: dst.alignment = copy(src.alignment)
                if src.border: dst.border = copy(src.border)
                if src.fill: dst.fill = copy(src.fill)
                if src.font: dst.font = copy(src.font)
        for name, col in col_map.items():
            if name not in row.index: continue
            if name in ("RSD", "Bât", "Ent", "Porte", "NDJI (J ELS→PAL)") or "PRÉVI. J+5 ouvrés" in name: continue
            ws.cell(i, col).value = clean_scalar(row[name])
        ws.cell(i, 3).value = f'=IFERROR(VLOOKUP($A{i},\'List LGMT\'!$A:$D,4,0),"")'
        ws.cell(i, 4).value = f'=IFERROR(VLOOKUP($A{i},\'List LGMT\'!$A:$E,5,0),"")'
        ws.cell(i, 5).value = f'=IFERROR(VLOOKUP($A{i},\'List LGMT\'!$A:$F,6,0),"")'
        ws.cell(i, 6).value = f'=IFERROR(VLOOKUP($A{i},\'List LGMT\'!$A:$L,12,0),"")'
        for idx, cat in enumerate(WORK_CATEGORIES):
            base = 9 + idx * 4
            start_col_letter = ws.cell(i, base + 1).column_letter
            ws.cell(i, base + 2).value = f'=IFERROR(IF({start_col_letter}{i}<>"",WORKDAY({start_col_letter}{i},5),""),"")'
        ws.cell(i, 41).value = f'=IFERROR(IF(AND(H{i}<>"",G{i}<>""),H{i}-G{i},""),"")'

    folder = os.path.dirname(os.path.abspath(path))
    suffix = ext if ext in (".xlsx", ".xlsm", ".xltx", ".xltm") else ".xlsx"
    fd, tmp = tempfile.mkstemp(prefix="oph65_save_", suffix=suffix, dir=folder)
    os.close(fd)
    backup = path + ".bak"
    try:
        wb.save(tmp)
        if not zipfile.is_zipfile(tmp):
            raise ValueError("Le classeur temporaire généré est invalide ; le fichier original n'a pas été modifié.")
        try: shutil.copy2(path, backup)
        except Exception: pass
        os.replace(tmp, path)
    finally:
        try: wb.close()
        except Exception: pass
        if os.path.exists(tmp):
            try: os.remove(tmp)
            except Exception: pass


def update_workbook_bytes(df, original_bytes, filename):
    """Met à jour une copie du classeur XLSX/XLSM et renvoie ses octets (VBA préservé pour XLSM)."""
    if not original_bytes:
        raise FileNotFoundError("Aucun classeur source n'est enregistré. Importez d'abord le fichier Excel/XLSM.")
    suffix=os.path.splitext(filename or "")[1].lower()
    if suffix not in (".xlsx",".xlsm",".xltx",".xltm"): suffix=".xlsm"
    fd,path=tempfile.mkstemp(prefix="oph65_cloud_",suffix=suffix); os.close(fd)
    try:
        with open(path,"wb") as f: f.write(original_bytes)
        save_to_source_xlsm(df,path)
        with open(path,"rb") as f: return f.read()
    finally:
        for candidate in (path,path+".bak"):
            try:
                if os.path.exists(candidate): os.remove(candidate)
            except Exception: pass


def persist_tracking_and_excel(df):
    """Sauvegarde la base persistante puis synchronise le classeur Excel conservé dans Supabase."""
    if CLOUD_MODE and supabase_ready():
        db_upsert_tracking(df)
        original=st.session_state.get("source_workbook_bytes")
        filename=st.session_state.get("source_workbook_name") or st.session_state.get("filename") or "Suivi_Travaux_OPH65.xlsm"
        if not original:
            original, saved_name=db_load_excel_bytes()
            if original:
                st.session_state.source_workbook_bytes=original
                st.session_state.source_workbook_name=saved_name
                filename=saved_name
        if original:
            updated=update_workbook_bytes(df,original,filename)
            st.session_state.source_workbook_bytes=updated
            st.session_state.source_workbook_name=filename
            db_save_excel_bytes(updated,filename)
            return True
        return False
    if st.session_state.get("source_path") and st.session_state.get("auto_save"):
        save_to_source_xlsm(df,st.session_state.source_path)
        return True
    return False


def export_flat_excel(df):
    output = io.BytesIO()
    clean = df.drop(columns=[c for c in df.columns if c.startswith("Statut ") or c.startswith("_mailto_") or c == "Statut global"], errors="ignore")
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        clean.to_excel(writer, index=False, sheet_name="Suivi travaux")
    return output.getvalue()


def delayed_days(row):
    delays = []
    for work, cols in WORK_COLUMNS.items():
        if row.get(f"Statut {work}") == "En retard" and pd.notna(row.get(cols["planned"])):
            delays.append((date.today() - pd.Timestamp(row[cols["planned"]]).date()).days)
    return max(delays) if delays else 0


def synthese_dataframe(df):
    e = enrich(df)
    rows = []
    for _, r in e.iterrows():
        actual_ends = [pd.Timestamp(r[c["real"]]) for c in WORK_COLUMNS.values() if c["real"] in r and pd.notna(r[c["real"]])]
        relocation = max(actual_ends).date() if r["Statut global"] == "Terminé" and actual_ends else None
        rows.append({
            "N° lot": r.get("N° LOT", ""), "RSD": r.get("RSD", ""), "Bât": r.get("Bât", ""),
            "Ent": r.get("Ent", ""), "Porte": r.get("Porte", ""),
            "Date de fin de travaux initiale (PAL)": r.get("PAL", pd.NaT),
            "Statut global": r.get("Statut global", ""),
            "Délai de retard (jours)": delayed_days(r) if r.get("Statut global") == "En retard" else 0,
            "Relocation possible à compter du": relocation,
        })
    return pd.DataFrame(rows)


def export_synthese(df):
    output = io.BytesIO()
    e = enrich(df)
    s = synthese_dataframe(df)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        counts = e["Statut global"].value_counts()
        summary = [{"Indicateur": "Total logements suivis", "Valeur": len(e)}]
        summary += [{"Indicateur": x, "Valeur": int(counts.get(x, 0))} for x in STATUS_LIST]
        pd.DataFrame(summary).to_excel(writer, index=False, sheet_name="Synthèse")
        s.to_excel(writer, index=False, sheet_name="Logements")
    return output.getvalue()



def export_synthese_pdf(df):
    """Génère une synthèse PDF A4 paysage avec une présentation institutionnelle OPH65."""
    if not HAS_REPORTLAB:
        raise RuntimeError("Le module ReportLab n'est pas installé. Exécutez : pip install -r requirements.txt")

    enriched = enrich(df)
    synth = synthese_dataframe(df).copy()
    output = io.BytesIO()
    page_w, page_h = landscape(A4)
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10*mm,
        leftMargin=10*mm,
        topMargin=12*mm,
        bottomMargin=12*mm,
        title="Synthèse d'avancement des travaux - OPH65",
        author="Office Public de l'Habitat des Hautes-Pyrénées",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "OPH65Title", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=15, leading=17, textColor=colors.HexColor("#174F91"), spaceAfter=2*mm,
    )
    subtitle_style = ParagraphStyle(
        "OPH65Subtitle", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=10, leading=12, textColor=colors.HexColor("#303A46"),
    )
    small_style = ParagraphStyle(
        "OPH65Small", parent=styles["Normal"], fontSize=7.2, leading=8.5,
        textColor=colors.HexColor("#4B5563"),
    )
    head_style = ParagraphStyle(
        "OPH65Head", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=6.6, leading=7.5, alignment=TA_CENTER, textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "OPH65Cell", parent=styles["Normal"], fontSize=6.4, leading=7.4,
        alignment=TA_CENTER,
    )

    story = []
    header_items = []
    if os.path.exists(LOGO_PATH):
        try:
            header_items.append(Image(LOGO_PATH, width=31*mm, height=18*mm, kind="proportional"))
        except Exception:
            header_items.append("")
    else:
        header_items.append("")
    header_text = [
        Paragraph("OFFICE PUBLIC DE L'HABITAT DES HAUTES-PYRÉNÉES", title_style),
        Paragraph("Synthèse d'avancement des travaux", subtitle_style),
        Paragraph(f"Éditée le {date.today().strftime('%d/%m/%Y')}", small_style),
    ]
    header_table = Table([[header_items[0], header_text]], colWidths=[36*mm, 235*mm])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LINEBELOW", (0,0), (-1,-1), 1.5, colors.HexColor("#174F91")),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story += [header_table, Spacer(1, 4*mm)]

    counts = enriched["Statut global"].value_counts() if not enriched.empty else pd.Series(dtype=int)
    metrics = [
        ("Logements suivis", len(enriched)),
        ("Terminés", int(counts.get("Terminé", 0))),
        ("En cours", int(counts.get("En cours", 0))),
        ("En retard", int(counts.get("En retard", 0))),
        ("Non démarrés", int(counts.get("Non démarré", 0))),
    ]
    metric_cells = []
    for label, value in metrics:
        metric_cells.append(Paragraph(f"<b>{label}</b><br/><font size='12'>{value}</font>", ParagraphStyle(
            f"metric_{label}", parent=small_style, alignment=TA_CENTER, leading=13,
        )))
    metric_table = Table([metric_cells], colWidths=[54*mm]*5)
    metric_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#F2F6FA")),
        ("BOX", (0,0), (-1,-1), 0.6, colors.HexColor("#9AA9B8")),
        ("INNERGRID", (0,0), (-1,-1), 0.4, colors.HexColor("#C8D2DC")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story += [metric_table, Spacer(1, 4*mm)]

    columns = [
        "N° lot", "RSD", "Bât", "Ent", "Porte",
        "Date de fin de travaux initiale (PAL)", "Statut global",
        "Délai de retard (jours)", "Relocation possible à compter du",
    ]
    labels = [
        "N° LOT", "RSD", "BÂT", "ENTRÉE", "PORTE", "DATE FIN TRAVAUX<br/>INITIALE (PAL)",
        "STATUT GLOBAL", "RETARD<br/>(JOURS)", "RELOCATION POSSIBLE<br/>À COMPTER DU",
    ]
    table_data = [[Paragraph(x, head_style) for x in labels]]
    for _, row in synth.iterrows():
        vals = []
        for c in columns:
            v = row.get(c, "")
            if c in ("Date de fin de travaux initiale (PAL)", "Relocation possible à compter du"):
                try:
                    v = "" if pd.isna(v) else pd.Timestamp(v).strftime("%d/%m/%Y")
                except Exception:
                    v = ""
            elif pd.isna(v):
                v = ""
            vals.append(Paragraph(html.escape(str(v)), cell_style))
        table_data.append(vals)

    col_widths = [18*mm, 38*mm, 14*mm, 16*mm, 17*mm, 39*mm, 28*mm, 21*mm, 43*mm]
    data_table = Table(table_data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    styles_cmd = [
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#174F91")),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#9CA3AF")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 3),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
        ("LEFTPADDING", (0,0), (-1,-1), 2),
        ("RIGHTPADDING", (0,0), (-1,-1), 2),
    ]
    for i, (_, row) in enumerate(synth.iterrows(), start=1):
        if i % 2 == 0:
            styles_cmd.append(("BACKGROUND", (0,i), (-1,i), colors.HexColor("#F7F9FB")))
        status = str(row.get("Statut global", ""))
        if status == "En retard":
            styles_cmd.append(("BACKGROUND", (6,i), (7,i), colors.HexColor("#FDECEC")))
            styles_cmd.append(("TEXTCOLOR", (6,i), (7,i), colors.HexColor("#9B1C1C")))
        elif status == "Terminé":
            styles_cmd.append(("BACKGROUND", (6,i), (6,i), colors.HexColor("#E7F4E7")))
            styles_cmd.append(("TEXTCOLOR", (6,i), (6,i), colors.HexColor("#1F6B2A")))
    data_table.setStyle(TableStyle(styles_cmd))
    story.append(data_table)

    def _footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#174F91"))
        canvas.setLineWidth(0.5)
        canvas.line(10*mm, 8*mm, page_w-10*mm, 8*mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5B6570"))
        canvas.drawString(10*mm, 4.5*mm, "OPH65 - Suivi des travaux")
        canvas.drawRightString(page_w-10*mm, 4.5*mm, f"Page {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output.getvalue()


def send_email_with_attachment(server, port, sender, password, recipients, subject, body, attachment_bytes=None, attachment_name=None, use_ssl=True, auth_required=True, use_starttls=True):
    """Envoie réellement le message depuis l'interface, sans ouvrir Outlook/Gmail ou un client local."""
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = sender, recipients, subject
    msg.set_content(body)
    if attachment_bytes is not None:
        filename = attachment_name or "Synthese_OPH65.pdf"
        mime, _ = mimetypes.guess_type(filename)
        maintype, subtype = (mime.split("/", 1) if mime and "/" in mime else ("application", "octet-stream"))
        msg.add_attachment(attachment_bytes, maintype=maintype, subtype=subtype, filename=filename)
    if use_ssl:
        with smtplib.SMTP_SSL(server, int(port), context=ssl.create_default_context(), timeout=25) as smtp:
            if auth_required: smtp.login(sender, password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(server, int(port), timeout=25) as smtp:
            smtp.ehlo()
            if use_starttls:
                smtp.starttls(context=ssl.create_default_context()); smtp.ehlo()
            if auth_required: smtp.login(sender, password)
            smtp.send_message(msg)


def _normalize_email_addresses(value):
    """Retourne une liste d'adresses SMTP simples et valides.

    Accepte les séparateurs virgule, point-virgule et retour ligne, ainsi que
    les formats "Nom <adresse@domaine.fr>". Les en-têtes restent lisibles,
    mais l'enveloppe SMTP n'utilise que les adresses réelles.
    """
    from email.utils import getaddresses, parseaddr
    import re

    if value is None:
        return []
    raw = str(value).replace(";", ",").replace("\r", "\n")
    parts = [p.strip() for p in raw.replace("\n", ",").split(",") if p.strip()]
    parsed = getaddresses(parts)
    out = []
    for _name, addr in parsed:
        addr = (addr or "").strip().replace("mailto:", "")
        # garde-fou volontairement simple : Gmail validera ensuite le domaine.
        if re.fullmatch(r"[^@\s<>]+@[^@\s<>]+\.[^@\s<>]+", addr):
            if addr not in out:
                out.append(addr)
    # getaddresses peut ignorer une adresse simple atypiquement séparée ; second passage
    if not out:
        for token in parts:
            addr = parseaddr(token)[1].strip()
            if re.fullmatch(r"[^@\s<>]+@[^@\s<>]+\.[^@\s<>]+", addr):
                out.append(addr)
    return out


def send_email_with_inline_qr(server, port, sender, password, recipients, subject, body, qr_bytes, use_ssl=True, auth_required=True, use_starttls=True):
    """Envoie le BC en multipart/alternative avec QR visible et joint.

    Les adresses sont normalisées avant l'appel SMTP afin d'éviter les commandes
    RCPT TO invalides (notamment listes séparées par ';'), que Gmail refuse en 5.5.2.
    """
    sender_list = _normalize_email_addresses(sender)
    recipient_list = _normalize_email_addresses(recipients)
    if len(sender_list) != 1:
        raise ValueError("Adresse expéditeur SMTP invalide.")
    if not recipient_list:
        raise ValueError("Aucune adresse destinataire valide. Utilisez une adresse mail par entreprise ; plusieurs adresses peuvent être séparées par une virgule ou un point-virgule.")

    sender_addr = sender_list[0]
    msg = EmailMessage()
    msg["From"] = sender_addr
    msg["To"] = ", ".join(recipient_list)
    msg["Subject"] = str(subject or "").replace("\r", " ").replace("\n", " ").strip()
    plain_body = str(body or "")
    msg.set_content(plain_body, charset="utf-8")

    safe_html = html.escape(plain_body).replace("\n", "<br>")
    html_body = (
        "<html><body><div style='font-family:Arial,sans-serif;font-size:14px'>"
        + safe_html
        + "<br><br><b>QR code de validation de fin de prestation :</b><br>"
          "<img src='cid:oph65qr' width='220' height='220' alt='QR code BC'>"
          "<br><small>Si le QR code ne s'affiche pas, utilisez la pièce jointe QR_BC.png.</small>"
          "</div></body></html>"
    )
    msg.add_alternative(html_body, subtype="html", charset="utf-8")
    html_part = msg.get_payload()[-1]
    html_part.add_related(
        bytes(qr_bytes), maintype="image", subtype="png",
        cid="oph65qr", disposition="inline", filename="QR_BC.png"
    )
    # Ajoute aussi le QR en pièce jointe pour les clients mail qui bloquent les images inline.
    msg.add_attachment(bytes(qr_bytes), maintype="image", subtype="png", filename="QR_BC.png")

    if use_ssl:
        with smtplib.SMTP_SSL(server, int(port), context=ssl.create_default_context(), timeout=25) as smtp:
            if auth_required:
                smtp.login(sender_addr, password)
            smtp.send_message(msg, from_addr=sender_addr, to_addrs=recipient_list)
    else:
        with smtplib.SMTP(server, int(port), timeout=25) as smtp:
            smtp.ehlo()
            if use_starttls:
                smtp.starttls(context=ssl.create_default_context())
                smtp.ehlo()
            if auth_required:
                smtp.login(sender_addr, password)
            smtp.send_message(msg, from_addr=sender_addr, to_addrs=recipient_list)


def generate_qr_png(url):
    if not HAS_QRCODE: raise RuntimeError("Le module qrcode n'est pas installé.")
    img=qrcode.make(url); out=io.BytesIO(); img.save(out,format="PNG"); return out.getvalue()


def create_bc_validation(lot_no, bc_no, work, company_email):
    token=secrets.token_urlsafe(24)
    client=get_supabase_client()
    if not client: raise RuntimeError("Supabase n'est pas configuré.")
    client.table("bc_validations").insert({
        "token":token,"bc_no":str(bc_no),"lot_no":normalize_lot_key(lot_no),"work_category":work,
        "company_email":company_email or "","status":"pending","created_at":datetime.utcnow().isoformat()
    }).execute()
    return token


def load_bc_validation(token):
    client=get_supabase_client()
    if not client:return None
    try:
        res=client.table("bc_validations").select("*").eq("token",token).limit(1).execute()
        rows=getattr(res,"data",None) or []
        return rows[0] if rows else None
    except Exception:return None


def _normalized_header_key(value):
    """Normalise un nom de colonne pour retrouver les variantes accentuées/non accentuées."""
    import unicodedata
    txt=unicodedata.normalize("NFKD", str(value or ""))
    txt="".join(ch for ch in txt if not unicodedata.combining(ch))
    return " ".join(txt.upper().replace("\n"," ").split())


def _resolve_tracking_column(df, expected):
    if expected in df.columns:
        return expected
    target=_normalized_header_key(expected)
    for col in df.columns:
        if _normalized_header_key(col)==target:
            return col
    return None


def complete_bc_validation(token):
    """Valide un BC, met à jour Supabase ET le classeur Excel synchronisé.

    La recherche s'effectue sur le triplet N° LOT + corps d'état + N° BC afin
    d'éviter toute mise à jour d'une mauvaise prestation.
    """
    validation=load_bc_validation(token)
    if not validation:
        raise ValueError("QR code inconnu ou expiré.")
    if validation.get("status")=="validated":
        return validation

    df=db_load_tracking()
    lot=normalize_lot_key(validation.get("lot_no"))
    work=str(validation.get("work_category") or "")
    bc_expected=str(validation.get("bc_no") or "").strip()
    if df.empty or "N° LOT" not in df.columns:
        raise ValueError("Le suivi persistant est vide ou indisponible.")

    cols=WORK_COLUMNS.get(work)
    if not cols:
        raise ValueError(f"Corps d'état inconnu : {work}")
    order_col=_resolve_tracking_column(df, cols.get("order"))
    real_col=_resolve_tracking_column(df, cols.get("real"))
    if not real_col:
        raise ValueError(f"La colonne de date réelle n'a pas été trouvée pour {work}.")

    lot_mask=df["N° LOT"].map(normalize_lot_key)==lot
    candidate_indexes=df.index[lot_mask].tolist()
    if not candidate_indexes:
        raise ValueError("Le logement correspondant n'existe plus dans le suivi.")

    # Si un N° BC est disponible, on exige qu'il corresponde à la prestation validée.
    idx=None
    for candidate in candidate_indexes:
        current_bc=""
        if order_col:
            value=df.at[candidate,order_col]
            current_bc="" if pd.isna(value) else str(value).strip()
        if not bc_expected or current_bc==bc_expected:
            idx=candidate
            break
    if idx is None:
        raise ValueError(f"Le BC {bc_expected} ne correspond plus au logement {lot} pour {work}.")

    validation_date=pd.Timestamp(date.today())
    df.at[idx,real_col]=validation_date
    df=normalize_date_dtypes(calculate_planned_dates(df))

    # 1) Base persistante Supabase.
    db_upsert_tracking(df)

    # Contrôle de lecture après écriture : on ne confirme jamais la validation
    # si la date réelle n'est pas effectivement persistée.
    check=db_load_tracking()
    check_rows=check.index[check["N° LOT"].map(normalize_lot_key)==lot].tolist() if not check.empty and "N° LOT" in check.columns else []
    if not check_rows:
        raise RuntimeError("La mise à jour Supabase n'a pas pu être vérifiée.")
    check_real=_resolve_tracking_column(check, real_col)
    saved_value=check.at[check_rows[0],check_real] if check_real else pd.NaT
    if pd.isna(saved_value) or pd.Timestamp(saved_value).date()!=validation_date.date():
        raise RuntimeError("La date de fin n'a pas été enregistrée dans Supabase.")

    # 2) Copie Excel persistante.
    original,name=db_load_excel_bytes()
    if original:
        updated=update_workbook_bytes(df,original,name)
        db_save_excel_bytes(updated,name)

    # 3) Historique de validation QR.
    validated_at=datetime.utcnow().isoformat()
    client=get_supabase_client()
    client.table("bc_validations").update({"status":"validated","validated_at":validated_at}).eq("token",token).execute()
    validation["status"]="validated"
    validation["validated_at"]=validated_at
    validation["real_date"]=validation_date.strftime("%Y-%m-%d")
    return validation


def render_public_bc_validation(token):
    validation=load_bc_validation(token)
    st.markdown("<style>[data-testid='stSidebar']{display:none}</style>",unsafe_allow_html=True)
    if os.path.exists(LOGO_PATH): st.image(LOGO_PATH,width=180)
    st.title("Validation de fin de prestation")
    if not validation:
        st.error("Ce QR code n'est pas valide."); st.stop()
    st.markdown(f"**Bon de commande :** {html.escape(str(validation.get('bc_no','')))}  \n**N° lot :** {html.escape(str(validation.get('lot_no','')))}  \n**Prestation :** {html.escape(WORK_DISPLAY_NAMES.get(validation.get('work_category'),validation.get('work_category','')))}")
    if validation.get("status")=="validated":
        st.success("La fin de prestation a déjà été validée. Merci."); st.stop()
    st.info("Après avoir terminé la prestation, validez ci-dessous. Le Responsable Technique de Secteur sera automatiquement informé.")
    if st.button("✅ Valider la fin de prestation",type="primary",use_container_width=True):
        try:
            val=complete_bc_validation(token)
            recipient=str(_secret("app.rts_email","") or load_config().get("rts_email","")).strip()
            if recipient:
                body=f"Bonjour,\n\nLa fin de prestation vient d'être validée par QR code.\n\nBC : {val.get('bc_no','')}\nN° lot : {val.get('lot_no','')}\nPrestation : {WORK_DISPLAY_NAMES.get(val.get('work_category'),val.get('work_category',''))}\nDate : {datetime.now().strftime('%d/%m/%Y %H:%M')}\n\nCordialement,\nOPH65"
                send_email_with_attachment(st.session_state.smtp_server,st.session_state.smtp_port,st.session_state.smtp_email,st.session_state.smtp_password,recipient,f"Fin de prestation – BC {val.get('bc_no','')}",body,use_ssl=st.session_state.smtp_use_ssl,auth_required=st.session_state.smtp_auth_required,use_starttls=st.session_state.smtp_use_starttls)
                db_log_mail("validation_bc",recipient,f"Fin de prestation – BC {val.get('bc_no','')}",val.get('lot_no',''),val.get('bc_no',''),val.get('work_category',''))
            st.success("Fin de prestation validée. Le Responsable Technique de Secteur a été informé."); st.balloons()
        except Exception as e: st.error(f"Impossible de valider : {e}")
    st.stop()


def test_smtp_connection(server, port, sender, password, use_ssl=True, auth_required=True, use_starttls=True):
    if use_ssl:
        with smtplib.SMTP_SSL(server, int(port), context=ssl.create_default_context(), timeout=20) as smtp:
            if auth_required: smtp.login(sender, password)
            smtp.noop()
    else:
        with smtplib.SMTP(server, int(port), timeout=20) as smtp:
            smtp.ehlo()
            if use_starttls:
                smtp.starttls(context=ssl.create_default_context()); smtp.ehlo()
            if auth_required: smtp.login(sender, password)
            smtp.noop()


def smtp_auth_error_message(exc):
    """Message utilisateur compréhensible pour les erreurs d'authentification SMTP."""
    return (
        "Authentification SMTP refusée. Vérifiez l'adresse expéditeur et le secret SMTP. "
        "Avec Gmail / Google Workspace, utilisez un mot de passe d'application si cette fonction est autorisée. "
        "Sinon, configurez le relais SMTP communiqué par le service informatique OPH65."
    )


@st.dialog("Enregistrement effectué")
def show_save_dialog():
    """Confirmation modale affichée après sauvegarde et remise à zéro du formulaire."""
    st.success(st.session_state.get("save_dialog_message", "Le logement a été enregistré avec succès."))
    st.write("Le formulaire a été réinitialisé et est prêt pour une nouvelle saisie.")
    if st.button("OK", type="primary", key="save_dialog_ok", use_container_width=True):
        st.rerun()

def format_date_fr(v):
    if v is None or pd.isna(v):
        return ""
    return pd.Timestamp(v).strftime("%d/%m/%Y")


def build_delay_email(work, row):
    cols = WORK_COLUMNS[work]
    order = row.get(cols["order"], "")
    order = "" if pd.isna(order) else str(order).strip()
    values = {
        "lot": work, "bc": order or "[à compléter]",
        "rsd": "" if pd.isna(row.get("RSD", "")) else str(row.get("RSD", "")),
        "bat": "" if pd.isna(row.get("Bât", "")) else str(row.get("Bât", "")),
        "ent": "" if pd.isna(row.get("Ent", "")) else str(row.get("Ent", "")),
        "porte": "" if pd.isna(row.get("Porte", "")) else str(row.get("Porte", "")),
        "echeance": format_date_fr(row.get(cols["planned"], pd.NaT)),
        "date": date.today().strftime("%d/%m/%Y"),
    }
    subject = st.session_state.delay_email_subject_template
    body = st.session_state.delay_email_body_template
    try: subject = subject.format(**values)
    except Exception: pass
    try: body = body.format(**values)
    except Exception: pass
    return st.session_state.lot_emails.get(work, ""), subject, body


def set_delay_mail_draft(row_index, work):
    """Prépare un nouveau brouillon à partir de la ligne et du corps d'état réellement cliqués."""
    try:
        enriched = enrich(st.session_state.df)
        # L'index transmis par AgGrid est stocké comme texte : on retrouve la ligne sans
        # supposer que l'index pandas est obligatoirement un entier.
        matches = [idx for idx in enriched.index if str(idx) == str(row_index)]
        if not matches:
            return False
        row = enriched.loc[matches[0]]
    except Exception:
        return False
    recipient, subject, body = build_delay_email(work, row)
    order_col = WORK_COLUMNS[work]["order"]
    bc = row.get(order_col, "")
    bc = "" if pd.isna(bc) else str(bc).strip()
    st.session_state.delay_mail_draft = {
        "row_index": str(row_index), "work": work, "bc": bc,
        "recipient": recipient, "subject": subject, "body": body,
        "label": f"{row.get('RSD','')} / bât. {row.get('Bât','')} / entrée {row.get('Ent','')} / porte {row.get('Porte','')}",
    }
    return True


def render_delay_mail_draft():
    """Affiche le brouillon de retard sous le tableau, modifiable puis envoyable."""
    draft = st.session_state.get("delay_mail_draft")
    if not draft:
        return
    st.divider()
    st.markdown("### 📧 Mail de retard prêt à envoyer")
    bc_label = draft.get("bc", "") or "non renseigné"
    st.caption(f"Lot : {draft.get('work','')} — BC : {bc_label} — logement {draft.get('label','')}")
    recipient = st.text_input("Destinataire", value=draft.get("recipient", ""), key="delay_inline_recipient")
    subject = st.text_input("Objet", value=draft.get("subject", ""), key="delay_inline_subject")
    body = st.text_area("Message", value=draft.get("body", ""), height=300, key="delay_inline_body")
    cclear, csend = st.columns([1, 1])
    if cclear.button("✖ Fermer le brouillon", use_container_width=True, key="delay_inline_close"):
        st.session_state.delay_mail_draft = None
        for k in ["delay_inline_recipient", "delay_inline_subject", "delay_inline_body"]:
            st.session_state.pop(k, None)
        st.session_state.delay_grid_version += 1
        st.rerun()
    if csend.button("📤 Envoyer le mail", type="primary", use_container_width=True, key="delay_inline_send"):
        if not recipient:
            st.error("Renseignez un destinataire.")
        elif not (st.session_state.smtp_server and st.session_state.smtp_email):
            st.error("Configurez d'abord le serveur SMTP dans Paramètres > Emails / SMTP.")
        elif st.session_state.smtp_auth_required and not st.session_state.smtp_password:
            st.error("Renseignez le mot de passe SMTP dans Paramètres > Emails / SMTP.")
        else:
            try:
                send_email_with_attachment(
                    st.session_state.smtp_server, st.session_state.smtp_port, st.session_state.smtp_email,
                    st.session_state.smtp_password, recipient, subject, body,
                    use_ssl=st.session_state.smtp_use_ssl,
                    auth_required=st.session_state.smtp_auth_required,
                    use_starttls=st.session_state.smtp_use_starttls,
                )
                # Après un envoi réussi, aucune donnée du mail précédent ne doit rester
                # active : le prochain clic « En retard » doit construire un brouillon neuf.
                st.session_state.delay_mail_draft = None
                for k in ["delay_inline_recipient", "delay_inline_subject", "delay_inline_body"]:
                    st.session_state.pop(k, None)
                st.session_state.delay_grid_version += 1
                st.session_state.delay_mail_sent_message = f"Mail envoyé à {recipient}. Le brouillon a été réinitialisé."
                st.rerun()
            except smtplib.SMTPAuthenticationError as e:
                st.error(smtp_auth_error_message(e))
            except Exception as e:
                st.error(f"Échec de l'envoi : {e}")


def _grid_date_column(name, series):
    """Détecte les colonnes qui doivent être affichées comme dates dans AgGrid."""
    n = str(name).upper()
    if _looks_like_date_column(str(name)) or "RELOCATION" in n or "ÉCHÉANCE" in n or "ECHEANCE" in n:
        return True
    return pd.api.types.is_datetime64_any_dtype(series)


def format_grid_dates(df):
    """Prépare les valeurs pour AgGrid : JJ/MM/AAAA au lieu des timestamps en ms et jamais de NaT."""
    out = df.copy()
    for c in out.columns:
        if str(c).startswith("_mailto_"):
            continue
        if _grid_date_column(c, out[c]):
            parsed = pd.to_datetime(out[c], errors="coerce")
            out[c] = parsed.dt.strftime("%d/%m/%Y").fillna("")
    # Évite l'affichage littéral de NaN / NaT dans les autres colonnes.
    out = out.astype(object).where(pd.notna(out), "")
    return out


def _with_interface_references(df):
    """Ajoute en tête les références logement à tous les tableaux de l'interface quand elles sont disponibles."""
    out = apply_reference_database(df)
    source = enrich(st.session_state.df) if not st.session_state.df.empty else pd.DataFrame()
    # Les tableaux issus du suivi conservent généralement l'index du logement : on peut donc
    # réinjecter les références sans modifier les données métier du tableau d'origine.
    if not source.empty:
        for ref in ["N° LOT", "RSD", "Bât", "Ent", "Porte"]:
            if ref not in out.columns and ref in source.columns:
                try:
                    out.insert(len([c for c in ["N° LOT", "RSD", "Bât", "Ent", "Porte"] if c in out.columns]), ref, source.reindex(out.index)[ref].values)
                except Exception:
                    pass
    refs = [c for c in ["N° LOT", "RSD", "Bât", "Ent", "Porte"] if c in out.columns]
    others = [c for c in out.columns if c not in refs]
    return out[refs + others]


def prepare_grid(df):
    grid = _with_interface_references(df)
    grid["_mail_row_index"] = [str(i) for i in grid.index]
    grid["_clicked_action"] = ""
    return format_grid_dates(grid)


def printable_table_component(df, title="Tableau de suivi OPH65", key="print"):
    """Bouton d'impression avec un modèle OPH65 indépendant de l'interface Streamlit."""
    printable = format_grid_dates(_with_interface_references(df))
    printable = printable.drop(columns=[c for c in printable.columns if str(c).startswith("_mail") or str(c).startswith("_mailto_")], errors="ignore")
    headers = [display_header(c).replace("\n", "<br>") for c in printable.columns]
    head_html = "".join(f"<th>{h}</th>" for h in headers)
    rows_html = []
    for _, row in printable.iterrows():
        cells = "".join(f"<td>{html.escape(str(row.get(c, '')))}</td>" for c in printable.columns)
        rows_html.append(f"<tr>{cells}</tr>")
    logo_html = ""
    if os.path.exists(LOGO_PATH):
        try:
            with open(LOGO_PATH, "rb") as f:
                logo64 = base64.b64encode(f.read()).decode("ascii")
            logo_html = f'<img src="data:image/jpeg;base64,{logo64}" class="logo">'
        except Exception:
            pass
    report = f"""
    <div class="toolbar"><button onclick="window.print()">🖨️ Imprimer ce tableau</button></div>
    <div class="report">
      <header>{logo_html}<div><h1>OFFICE PUBLIC DE L'HABITAT DES HAUTES-PYRÉNÉES</h1><h2>{html.escape(title)}</h2><p>Édité le {date.today().strftime('%d/%m/%Y')} — {len(printable)} ligne(s)</p></div></header>
      <table><thead><tr>{head_html}</tr></thead><tbody>{''.join(rows_html)}</tbody></table>
      <footer>OPH 65 — Suivi des travaux</footer>
    </div>
    <style>
      body{{font-family:Arial,sans-serif;margin:0}} .toolbar{{padding:4px 0}}
      button{{background:#174f91;color:#fff;border:0;border-radius:7px;padding:9px 14px;font-weight:700;cursor:pointer}}
      .report{{display:none}}
      @media print{{
        @page{{size:A4 landscape;margin:9mm}} .toolbar{{display:none}} .report{{display:block}}
        header{{display:flex;align-items:center;border-bottom:3px solid #174f91;padding-bottom:8px;margin-bottom:10px}}
        .logo{{width:95px;max-height:62px;object-fit:contain;margin-right:18px}} h1{{font-size:15px;color:#174f91;margin:0 0 4px}} h2{{font-size:13px;margin:0 0 3px}} p{{font-size:9px;margin:0;color:#555}}
        table{{width:100%;border-collapse:collapse;font-size:7px;table-layout:auto}} th{{background:#174f91;color:white;font-weight:700;white-space:normal;line-height:1.15}}
        th,td{{border:1px solid #9ca3af;padding:3px 4px;text-align:center;vertical-align:middle;word-break:normal}} tr{{break-inside:avoid}} thead{{display:table-header-group}}
        footer{{font-size:8px;margin-top:7px;text-align:right;color:#555}}
      }}
    </style>
    """
    components.html(report, height=48, scrolling=False)


def show_excel_grid(df, height=560, key="grid", delay_click_work=None):
    if df.empty:
        st.info("Aucune ligne à afficher.")
        return
    if not HAS_AGGRID:
        st.warning("Le module de grille avancée n'est pas installé. Installez les dépendances avec `pip install -r requirements.txt` pour activer les filtres avancés.")
        fallback_src = _with_interface_references(df)
        fallback = fallback_src.rename(columns={c: display_header(c) for c in fallback_src.columns})
        st.dataframe(fallback, use_container_width=True, height=height)
        printable_table_component(fallback_src, title="Tableau de suivi OPH65", key=f"print_{key}")
        return
    grid = prepare_grid(df)
    gb = GridOptionsBuilder.from_dataframe(grid)
    gb.configure_default_column(sortable=True, filter=True, floatingFilter=True, resizable=True, minWidth=95, wrapHeaderText=True, autoHeaderHeight=True)
    for c in grid.columns:
        if c.startswith("_mailto_") or c in ("_mail_row_index", "_clicked_action"):
            gb.configure_column(c, hide=True)
        else:
            gb.configure_column(c, headerName=display_header(c))
    # Style : ligne entièrement verte si terminée ; références rouges si un retard ; statuts selon leur couleur.
    status_colors = st.session_state.status_colors
    style_js = JsCode(f"""
    function(params) {{
      const field = params.colDef.field || '';
      const global = params.data ? params.data['Statut global'] : '';
      if (field === 'N° LOT') return {{textDecoration:'underline', cursor:'pointer', fontWeight:'700', color:'#174f91'}};
      if (global === 'Terminé') return {{backgroundColor:'{status_colors['Terminé']}', color:'#1b5e20'}};
      if (global === 'En retard' && ['N° LOT','GDC','RSD','Bât','Ent','Porte'].includes(field))
        return {{backgroundColor:'{status_colors['En retard']}', color:'#8a1c1c', fontWeight:'600'}};
      if (field.startsWith('Statut ')) {{
        if (params.value === 'En retard') return {{backgroundColor:'{status_colors['En retard']}', color:'#8a1c1c', fontWeight:'700', textDecoration:'underline', cursor:'pointer'}};
        if (params.value === 'Terminé') return {{backgroundColor:'{status_colors['Terminé']}', color:'#1b5e20'}};
        if (params.value === 'En cours') return {{backgroundColor:'{status_colors['En cours']}', color:'#0b3d75'}};
        if (params.value === 'Non démarré') return {{backgroundColor:'{status_colors['Non démarré']}', color:'#444'}};
      }}
      return null;
    }}
    """)
    for c in grid.columns:
        if not c.startswith("_mailto_") and c not in ("_mail_row_index", "_clicked_action"):
            gb.configure_column(c, cellStyle=style_js)
    # Le N° LOT est cliquable dans toutes les grilles. En retard conserve en plus son action mail.
    gb.configure_selection(selection_mode="single", use_checkbox=False, suppressRowClickSelection=True)
    delay_field = f"Statut {delay_click_work}" if delay_click_work else ""
    click_js = JsCode(f"""
    function(params) {{
      if (!params || !params.colDef || !params.data || !params.node) return;
      const field = params.colDef.field || '';
      if (field === 'N° LOT' && params.value) {{
        params.node.setDataValue('_clicked_action', 'identity');
        params.api.deselectAll();
        params.node.setSelected(true, true);
        return;
      }}
      if ({json.dumps(bool(delay_click_work))} && field === {json.dumps(delay_field, ensure_ascii=False)} && params.value === 'En retard') {{
        params.node.setDataValue('_clicked_action', 'delay');
        params.api.deselectAll();
        params.node.setSelected(true, true);
      }}
    }}
    """)
    gb.configure_grid_options(onCellClicked=click_js)
    response = AgGrid(
        grid, gridOptions=gb.build(), allow_unsafe_jscode=True, theme="streamlit", height=height,
        fit_columns_on_grid_load=False, key=(f"{key}_v{st.session_state.delay_grid_version}" if delay_click_work else key),
        update_on=["selectionChanged"]
    )
    selected = getattr(response, "selected_rows", None)
    if selected is None and isinstance(response, dict):
        selected = response.get("selected_rows")
    try:
        if isinstance(selected, pd.DataFrame): records = selected.to_dict("records")
        elif isinstance(selected, dict): records = [selected]
        else: records = list(selected) if selected is not None else []
        if records:
            rec = records[0]
            action = rec.get("_clicked_action", "")
            if action == "identity" and rec.get("N° LOT"):
                st.session_state.identity_lot = normalize_lot_key(rec.get("N° LOT"))
                st.session_state.identity_return_page = st.session_state.get("navigation_page", "🏠 Tableau de bord")
                st.rerun()
            elif action == "delay" and delay_click_work:
                row_idx = rec.get("_mail_row_index")
                if row_idx is not None:
                    current = st.session_state.get("delay_mail_draft") or {}
                    if current.get("row_index") != str(row_idx) or current.get("work") != delay_click_work:
                        for wk in ["delay_inline_recipient", "delay_inline_subject", "delay_inline_body"]:
                            st.session_state.pop(wk, None)
                        set_delay_mail_draft(row_idx, delay_click_work)
    except Exception:
        pass
    printable_table_component(df, title="Tableau de suivi OPH65", key=f"print_{key}")


def chart_detail_table(df, work, status):
    sc = f"Statut {work}"
    detail = df[df[sc] == status].copy()
    cols = [c for c in REFERENCE_COLUMNS + [WORK_COLUMNS[work]["order"], WORK_COLUMNS[work]["start"], WORK_COLUMNS[work]["planned"], WORK_COLUMNS[work]["real"], sc, "Statut global"] if c in detail.columns]
    st.markdown(f"#### Logements — {work} / {status}")
    if status == "En retard":
        st.caption("Cliquez directement sur une cellule rouge « En retard » pour préparer le mail correspondant sous le tableau.")
    show_excel_grid(detail[cols], height=360, key=f"dash_detail_{work}_{status}", delay_click_work=work if status == "En retard" else None)
    if status == "En retard":
        if st.session_state.get("delay_mail_sent_message"):
            st.success(st.session_state.delay_mail_sent_message)
            st.session_state.delay_mail_sent_message = ""
        render_delay_mail_draft()



def bc_management_dataframe(df):
    e=enrich(df); rows=[]
    for idx,r in e.iterrows():
        for work,cols in WORK_COLUMNS.items():
            bc=r.get(cols["order"],""); bc="" if pd.isna(bc) else str(bc).strip()
            if not bc: continue
            status=r.get(f"Statut {work}","Non démarré")
            if status not in ("En cours","En retard"): continue
            rows.append({
                "N° LOT":r.get("N° LOT",""),"RSD":r.get("RSD",""),"Bât":r.get("Bât",""),"Ent":r.get("Ent",""),"Porte":r.get("Porte",""),
                "Corps d'état":work,"N° BC":bc,"Date BC":r.get(cols["start"],pd.NaT),"Date théorique":r.get(cols["planned"],pd.NaT),
                "Statut":status,"Entreprise / mail":st.session_state.lot_emails.get(work,""),"_row_index":str(idx)
            })
    out=pd.DataFrame(rows)
    if not out.empty: out=out.sort_values(["N° LOT","Statut","Corps d'état"],kind="stable")
    return out


def prepare_bc_mail(record):
    bc=str(record.get("N° BC","")).strip(); lot=normalize_lot_key(record.get("N° LOT")); work=record.get("Corps d'état","")
    recipient=st.session_state.lot_emails.get(work,"") or str(record.get("Entreprise / mail","") or "")
    token=create_bc_validation(lot,bc,work,recipient)
    base=str(_secret("app.public_base_url","") or "").rstrip("/")
    if not base: raise RuntimeError("Renseignez app.public_base_url dans les Secrets Streamlit.")
    url=f"{base}/?bc_token={urllib.parse.quote(token)}"
    qr=generate_qr_png(url)
    body=("Ce bon de commande est à réaliser, en cours de réalisation ou en retard. Merci de faire le nécessaire pour la réalisation des travaux dans le temps imparti conformément au CCTP.\n\n"
          "A l'issue, merci de scanner le QR code joint et valider l'envoi du mail pour informer le Responsable Technique de Secteur de la fin de votre prestation.")
    st.session_state.bc_mail_draft={"recipient":recipient,"subject":f"BC {bc}","body":body,"qr":qr,"url":url,"bc":bc,"lot":lot,"work":work}


def render_bc_mail_draft():
    d=st.session_state.get("bc_mail_draft")
    if not d:return
    st.markdown("### ✉️ Mail du bon de commande")
    st.caption(f"Lot {d.get('lot')} — {WORK_DISPLAY_NAMES.get(d.get('work'),d.get('work'))} — BC {d.get('bc')}")
    recipient=st.text_input("Destinataire",value=d.get("recipient",""),key="bc_recipient", help="Adresse mail de l’entreprise. Plusieurs adresses : séparez-les par une virgule ou un point-virgule.")
    subject=st.text_input("Objet",value=d.get("subject",""),key="bc_subject")
    body=st.text_area("Message",value=d.get("body",""),height=220,key="bc_body")
    if d.get("qr"): st.image(d["qr"],caption="QR code de validation",width=180)
    c1,c2=st.columns(2)
    if c1.button("✖ Fermer",use_container_width=True,key="bc_close"):
        st.session_state.bc_mail_draft=None
        for k in ("bc_recipient","bc_subject","bc_body"):st.session_state.pop(k,None)
        st.rerun()
    if c2.button("📤 Envoyer le BC",type="primary",use_container_width=True,key="bc_send"):
        try:
            send_email_with_inline_qr(st.session_state.smtp_server,st.session_state.smtp_port,st.session_state.smtp_email,st.session_state.smtp_password,recipient,subject,body,d["qr"],use_ssl=st.session_state.smtp_use_ssl,auth_required=st.session_state.smtp_auth_required,use_starttls=st.session_state.smtp_use_starttls)
            db_log_mail("bc",recipient,subject,d.get("lot"),d.get("bc"),d.get("work"))
            st.session_state.bc_mail_draft=None
            for k in ("bc_recipient","bc_subject","bc_body"):st.session_state.pop(k,None)
            st.success(f"BC {d.get('bc')} envoyé à {recipient}.")
        except smtplib.SMTPRecipientsRefused as e:
            st.error("Gmail a refusé le destinataire. Vérifiez l’adresse mail de l’entreprise dans Paramètres → Destinataires par lot.")
        except smtplib.SMTPDataError as e:
            st.error(f"Gmail a refusé le contenu du message ({getattr(e, 'smtp_code', '')}). Vérifiez les adresses puis réessayez. Détail : {e}")
        except Exception as e: st.error(f"Échec de l'envoi : {e}")


def render_bc_management():
    st.title("📦 Gestion des BC")
    if st.session_state.df.empty:
        st.warning("Aucune donnée de suivi disponible."); return
    bcdf=bc_management_dataframe(st.session_state.df)
    if bcdf.empty:
        st.success("Aucun bon de commande en cours ou en retard."); return
    a,b,c=st.columns(3); a.metric("BC actifs",len(bcdf)); b.metric("En retard",int((bcdf["Statut"]=="En retard").sum())); c.metric("En cours",int((bcdf["Statut"]=="En cours").sum()))
    status_filter=st.multiselect("Statut",["En retard","En cours"],default=["En retard","En cours"]); view=bcdf[bcdf["Statut"].isin(status_filter)].copy()
    display=view.drop(columns=["_row_index"],errors="ignore")
    show_excel_grid(display,height=450,key="bc_management_grid")
    choices=[]; mapping={}
    for i,r in view.reset_index(drop=True).iterrows():
        work_name = WORK_DISPLAY_NAMES.get(r["Corps d'état"], r["Corps d'état"])
        label=f"Lot {r['N° LOT']} — BC {r['N° BC']} — {work_name} — {r['Statut']}"
        choices.append(label); mapping[label]=r.to_dict()
    selected=st.selectbox("BC à envoyer / renvoyer",choices,key="bc_choice")
    if st.button("✉️ Préparer le mail et le QR code",type="primary",use_container_width=True):
        try:
            for k in ("bc_recipient","bc_subject","bc_body"):st.session_state.pop(k,None)
            prepare_bc_mail(mapping[selected]); st.rerun()
        except Exception as e: st.error(str(e))
    render_bc_mail_draft()


def _identity_value(value, default="Non renseigné"):
    text = _clean(value)
    return text if text else default


def _heating_label(code):
    code = _clean(code).upper()
    labels = {
        "GI": "Gaz individuel", "GC": "Gaz collectif", "EL": "Électrique",
        "GPI": "Gaz propane individuel", "BL": "Bois de chauffage", "RAE": "Radiateur à eau",
    }
    return f"{code} — {labels.get(code, 'Code non référencé')}" if code else "Non renseigné"


def _pes_label(value):
    text = _clean(value).strip()
    if text in ("1", "1.0", "Oui", "OUI", "oui", "True", "true"):
        return "Oui"
    if text in ("0", "0.0", "Non", "NON", "non", "False", "false"):
        return "Non"
    return text or "Non renseigné"


def render_housing_identity(lot_no):
    key = normalize_lot_key(lot_no)
    ref = load_reference_database().get(key)
    if not ref:
        st.error(f"Le lot {key or lot_no} n'existe pas dans la base logements.")
        if st.button("← Retour au tableau", type="primary"):
            st.session_state.identity_lot = ""
            st.rerun()
        return
    service = _clean(ref.get("Date mise en service", ""))
    try:
        service = pd.to_datetime(service).strftime("%d/%m/%Y") if service else "Non renseigné"
    except Exception:
        service = service or "Non renseigné"
    if st.button("← Retour au tableau", type="primary", key="identity_back_top"):
        st.session_state.identity_lot = ""
        st.rerun()
    st.markdown(f"""
    <div style="border:1px solid #d7dee8;border-radius:18px;overflow:hidden;background:white;box-shadow:0 8px 26px rgba(15,23,42,.08);margin:8px 0 20px 0">
      <div style="background:linear-gradient(135deg,{st.session_state.primary_color},#315f8f);color:white;padding:22px 26px">
        <div style="font-size:13px;opacity:.9;letter-spacing:.08em;text-transform:uppercase">OPH65 · Fiche d'identité logement</div>
        <div style="font-size:29px;font-weight:800;margin-top:4px">Lot n° {html.escape(key)}</div>
        <div style="font-size:15px;margin-top:5px">{html.escape(_identity_value(ref.get('Résidence')))} · RSD {html.escape(_identity_value(ref.get('RSD')))}</div>
      </div>
      <div style="padding:22px 26px;color:{st.session_state.font_color};font-family:{st.session_state.font_family},sans-serif">
        <div style="display:grid;grid-template-columns:repeat(4,minmax(130px,1fr));gap:12px;margin-bottom:20px">
          <div class="oph-card"><b>Bâtiment</b><span>{html.escape(_identity_value(ref.get('Bât')))}</span></div>
          <div class="oph-card"><b>Entrée</b><span>{html.escape(_identity_value(ref.get('Ent')))}</span></div>
          <div class="oph-card"><b>Porte</b><span>{html.escape(_identity_value(ref.get('Porte')))}</span></div>
          <div class="oph-card"><b>Typologie</b><span>{html.escape(_identity_value(ref.get('Typologie')))}</span></div>
        </div>
        <div class="oph-section-title">Localisation</div>
        <div class="oph-identity-grid">
          <div><b>Adresse</b><span>{html.escape(_identity_value(ref.get('Adresse')))}</span></div>
          <div><b>Suite adresse</b><span>{html.escape(_identity_value(ref.get('Suite adresse')))}</span></div>
          <div><b>Code postal</b><span>{html.escape(_identity_value(ref.get('Code postal')))}</span></div>
          <div><b>Ville</b><span>{html.escape(_identity_value(ref.get('Ville')))}</span></div>
          <div><b>Secteur</b><span>{html.escape(_identity_value(ref.get('Secteur')))}</span></div>
          <div><b>Date de mise en service</b><span>{html.escape(service)}</span></div>
        </div>
        <div class="oph-section-title" style="margin-top:20px">Équipements & contrat</div>
        <div class="oph-identity-grid">
          <div><b>Type de chauffage</b><span>{html.escape(_heating_label(ref.get('Type chauffage')))}</span></div>
          <div><b>Contrat PES</b><span>{html.escape(_pes_label(ref.get('Contrat PES')))}</span></div>
        </div>
      </div>
    </div>
    <style>
      .oph-card{{background:#f3f5f7;border:1px solid #e1e6ec;border-radius:12px;padding:13px 15px;display:flex;flex-direction:column;gap:5px}}
      .oph-card b,.oph-identity-grid b{{font-size:11px;text-transform:uppercase;letter-spacing:.04em;color:#687386}}
      .oph-card span{{font-size:17px;font-weight:700;color:#172033}}
      .oph-section-title{{font-size:17px;font-weight:800;color:{st.session_state.primary_color};padding-bottom:7px;border-bottom:2px solid #e6ebf1}}
      .oph-identity-grid{{display:grid;grid-template-columns:repeat(2,minmax(220px,1fr));gap:0 28px}}
      .oph-identity-grid>div{{display:flex;justify-content:space-between;gap:20px;padding:12px 0;border-bottom:1px solid #edf0f3}}
      .oph-identity-grid span{{font-weight:600;text-align:right;color:#172033}}
      @media(max-width:900px){{.oph-identity-grid{{grid-template-columns:1fr}}}}
    </style>
    """, unsafe_allow_html=True)
    if st.button("← Retour au tableau", key="identity_back_bottom"):
        st.session_state.identity_lot = ""
        st.rerun()

cfg = load_config()
def ss_default(name, value):
    if name not in st.session_state: st.session_state[name] = value

ss_default("df", pd.DataFrame())
ss_default("filename", "")
ss_default("source_path", cfg.get("source_path", ""))
ss_default("auto_save", bool(cfg.get("auto_save", True)))
ss_default("primary_color", cfg.get("primary_color", DEFAULT_PRIMARY_COLOR))
ss_default("sidebar_color", cfg.get("sidebar_color", DEFAULT_SIDEBAR_COLOR))
ss_default("status_colors", cfg.get("status_colors", DEFAULT_STATUS_COLORS.copy()))
ss_default("username", cfg.get("username", "Administrateur"))
ss_default("users", cfg.get("users", [cfg.get("username", "Administrateur")]) or ["Administrateur"])
ss_default("date_format", cfg.get("date_format", "JJ/MM/AAAA (31/12/2025)"))
ss_default("residence_options", cfg.get("residences", []))
ss_default("lot_emails", cfg.get("lot_emails", {}))
ss_default("smtp_server", cfg.get("smtp_server", "smtp.gmail.com"))
ss_default("smtp_port", int(cfg.get("smtp_port", 465)))
ss_default("smtp_use_ssl", bool(cfg.get("smtp_use_ssl", True)))
ss_default("smtp_email", cfg.get("smtp_email", ""))
ss_default("smtp_password", load_secure_smtp_password(cfg.get("smtp_email", "")))
ss_default("email_recipients", cfg.get("email_recipients", ""))
ss_default("analysis_email_subject_template", cfg.get("analysis_email_subject_template", "OPH65 – Synthèse d'avancement des travaux – {date}"))
ss_default("analysis_email_body", cfg.get("analysis_email_body", "Bonjour,\n\nVeuillez trouver ci-joint la synthèse d'avancement des travaux de l'ensemble des logements suivis.\n\nCordialement,\n\nM. Cédric SCHMIT\nOffice Public de l'Habitat des Hautes-Pyrénées\nResponsable Technique de Secteur TN2"))
ss_default("delay_email_subject_template", cfg.get("delay_email_subject_template", "Retard travaux – lot {lot} – BC n° {bc}"))
ss_default("delay_email_body_template", cfg.get("delay_email_body_template", load_config().get("delay_email_body_template", "")))
ss_default("smtp_auth_required", bool(cfg.get("smtp_auth_required", True)))
ss_default("smtp_use_starttls", bool(cfg.get("smtp_use_starttls", not bool(cfg.get("smtp_use_ssl", True)))))
ss_default("delay_mail_draft", None)
ss_default("delay_grid_version", 0)
ss_default("delay_mail_sent_message", "")
ss_default("background_image", cfg.get("background_image", ""))
ss_default("background_overlay", float(cfg.get("background_overlay", 0.22)))
ss_default("background_mode", cfg.get("background_mode", "Couleur"))
ss_default("background_color", cfg.get("background_color", "#f5f7fb"))
ss_default("font_family", cfg.get("font_family", "Arial"))
ss_default("font_color", cfg.get("font_color", "#1f2937"))
ss_default("identity_lot", "")
ss_default("identity_return_page", "🏠 Tableau de bord")
ss_default("edit_form_version", 0)
ss_default("edit_form_reset_pending", False)
ss_default("save_dialog_pending", False)
ss_default("save_dialog_message", "Le logement a été enregistré avec succès.")
ss_default("source_workbook_bytes", None)
ss_default("source_workbook_name", "")
ss_default("bc_mail_draft", None)
ss_default("rts_email", cfg.get("rts_email", str(_secret("app.rts_email", "") or "")))

if "autoload_done" not in st.session_state:
    st.session_state.autoload_done = True
    if CLOUD_MODE and supabase_ready():
        try:
            cloud_df=db_load_tracking()
            if not cloud_df.empty:
                st.session_state.df=cloud_df
                wb_bytes,wb_name=db_load_excel_bytes()
                st.session_state.source_workbook_bytes=wb_bytes
                st.session_state.source_workbook_name=wb_name
                st.session_state.filename=wb_name or "Base Supabase"
        except Exception:
            pass
    elif st.session_state.source_path and os.path.exists(st.session_state.source_path):
        try:
            st.session_state.df = read_excel(st.session_state.source_path)
            st.session_state.filename = os.path.basename(st.session_state.source_path)
        except Exception:
            pass

page_background_css = background_css(st.session_state.background_image, st.session_state.background_overlay, st.session_state.background_mode, st.session_state.background_color)
st.markdown(f"""
<style>
:root {{ --oph:{st.session_state.primary_color}; }}
{page_background_css}
[data-testid="stSidebar"] {{ background:{st.session_state.sidebar_color}; }}
[data-testid="stSidebar"] * {{ color:white; }}
/* Zone utilisateur : fond gris clair et texte foncé, indépendamment du thème Streamlit. */
[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] > div,
[data-testid="stSidebar"] [data-testid="stSelectbox"] div[role="combobox"] {{
    background:#e5e7eb !important; border:1px solid #cbd5e1 !important; color:#1f2937 !important;
}}
[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] span,
[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] div,
[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] svg {{
    color:#1f2937 !important; fill:#1f2937 !important;
}}
[data-testid="stSidebar"] [data-testid="stSelectbox"] label p {{ color:white !important; }}
[data-testid="stAppViewContainer"] {{ color:{st.session_state.font_color}; font-family:{st.session_state.font_family}, sans-serif; }}
[data-testid="stAppViewContainer"] p, [data-testid="stAppViewContainer"] label,
[data-testid="stAppViewContainer"] input, [data-testid="stAppViewContainer"] textarea,
[data-testid="stAppViewContainer"] button {{ font-family:{st.session_state.font_family}, sans-serif; }}
h1,h2,h3 {{ color:{st.session_state.primary_color}; font-family:{st.session_state.font_family}, sans-serif; }}
[data-testid='stSidebar'] button[kind='primary'] {{ background:#c62828 !important; border-color:#a61f1f !important; color:white !important; font-weight:700 !important; }}
.small {{color:#667085;font-size:.9rem;}}
</style>
""", unsafe_allow_html=True)

# Les liens QR ouvrent une page publique minimale, sans accès au reste de l'application.
try:
    _bc_token = st.query_params.get("bc_token", "")
except Exception:
    _bc_token = ""
if _bc_token:
    render_public_bc_validation(str(_bc_token))

# Protection simple de l'interface interne sur une URL Streamlit publique.
_access_password = str(_secret("app.access_password", "") or "")
if CLOUD_MODE and _access_password and not st.session_state.get("oph65_authenticated", False):
    if os.path.exists(LOGO_PATH): st.image(LOGO_PATH, width=190)
    st.title("OPH65 — Accès sécurisé")
    entered=st.text_input("Mot de passe d'accès",type="password",key="cloud_access_password")
    if st.button("Se connecter",type="primary",use_container_width=True):
        if secrets.compare_digest(entered,_access_password):
            st.session_state.oph65_authenticated=True; st.rerun()
        else:
            st.error("Mot de passe incorrect.")
    st.stop()

# Sur le Cloud, Supabase est la source de vérité. Un QR code est validé dans une
# autre session navigateur : on recharge donc le suivi à chaque rerun de
# l'interface interne pour refléter immédiatement les validations externes.
if CLOUD_MODE and supabase_ready() and (not _access_password or st.session_state.get("oph65_authenticated", False)):
    try:
        _fresh_tracking=db_load_tracking()
        if not _fresh_tracking.empty:
            st.session_state.df=_fresh_tracking
            _wb,_wb_name=db_load_excel_bytes()
            if _wb:
                st.session_state.source_workbook_bytes=_wb
                st.session_state.source_workbook_name=_wb_name
    except Exception:
        pass

with st.sidebar:
    if os.path.exists(LOGO_PATH): st.image(LOGO_PATH, use_container_width=True)
    st.markdown("### 🏗️ Suivi des travaux")
    active_user = st.selectbox("👤 Utilisateur", st.session_state.users, index=st.session_state.users.index(st.session_state.username) if st.session_state.username in st.session_state.users else 0, key="sidebar_user")
    if active_user != st.session_state.username:
        st.session_state.username = active_user
        save_config(username=active_user, users=st.session_state.users)
        st.rerun()
    page = st.radio("Navigation", [
        "🏠 Tableau de bord", "📋 Suivi des logements", "🔎 Consultation des données",
        "➕ Ajouter / Modifier", "📦 Gestion des BC", "📊 Analyses", "📁 Import / Export", "⚙️ Paramètres"
    ], key="navigation_page")
    st.divider()
    if not CLOUD_MODE:
        if st.button("🚪 Quitter l’application", use_container_width=True, type="primary"):
            confirm_application_exit()
    else:
        st.caption("☁️ Mode Streamlit Cloud")
    st.caption("V2.16 Cloud — Supabase, Excel synchronisé & Gestion des BC")

if st.session_state.get("identity_lot"):
    render_housing_identity(st.session_state.identity_lot)
elif page == "📁 Import / Export":
    st.title("📁 Import / Export Excel")
    if CLOUD_MODE:
        st.info("☁️ Sur Streamlit Cloud, importez le classeur depuis votre ordinateur. Un chemin C:\\... local n'est pas accessible au serveur distant.")
        uploaded = st.file_uploader("Importer le fichier Excel / XLSM de suivi", type=["xlsx", "xlsm", "xls"], key="cloud_main_upload")
    else:
        tab_path, tab_upload = st.tabs(["📂 Ouvrir le fichier de travail", "⬆️ Import ponctuel"])
        with tab_path:
            path_input = st.text_input("Chemin complet du fichier Excel / XLSM", value=st.session_state.source_path, placeholder=r"C:\Dossier\Suivi_Travaux.xlsm")
            if st.button("📥 Charger ce fichier"):
                if not path_input: st.error("Merci de renseigner un chemin.")
                elif not os.path.exists(path_input): st.error("Fichier introuvable.")
                else:
                    try:
                        st.session_state.df = read_excel(path_input)
                        st.session_state.filename = os.path.basename(path_input)
                        st.session_state.source_path = path_input
                        save_config(source_path=path_input)
                        st.success(f"Fichier chargé : {st.session_state.filename} — {len(st.session_state.df)} logement(s).")
                    except Exception as e: st.error(f"Impossible de lire le fichier : {e}")
            new_auto_save = st.checkbox("Mettre à jour automatiquement le classeur d'origine", value=st.session_state.auto_save)
            if new_auto_save != st.session_state.auto_save:
                st.session_state.auto_save = new_auto_save
                save_config(auto_save=new_auto_save)
            if st.session_state.source_path:
                st.caption(f"📌 Fichier mémorisé : {st.session_state.source_path}")
        with tab_upload:
            uploaded = st.file_uploader("Importer un fichier Excel", type=["xlsx", "xlsm", "xls"])
    if uploaded is not None:
        try:
            uploaded_bytes=uploaded.getvalue()
            st.session_state.df = read_excel(io.BytesIO(uploaded_bytes))
            st.session_state.filename = uploaded.name
            st.session_state.source_path = ""
            st.session_state.source_workbook_bytes=uploaded_bytes
            st.session_state.source_workbook_name=uploaded.name
            if CLOUD_MODE and supabase_ready():
                db_upsert_tracking(st.session_state.df)
                db_save_excel_bytes(uploaded_bytes,uploaded.name)
            st.success(f"Fichier chargé et synchronisé : {uploaded.name} — {len(st.session_state.df)} logement(s).")
        except Exception as e: st.error(f"Impossible de lire le fichier : {e}")
    if not st.session_state.df.empty:
        if CLOUD_MODE and supabase_ready():
            latest_bytes=st.session_state.get("source_workbook_bytes")
            latest_name=st.session_state.get("source_workbook_name") or "Suivi_Travaux_OPH65.xlsm"
            if not latest_bytes:
                latest_bytes,latest_name=db_load_excel_bytes()
            if latest_bytes:
                mime="application/vnd.ms-excel.sheet.macroEnabled.12" if latest_name.lower().endswith(".xlsm") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                st.download_button("⬇️ Télécharger le classeur Excel synchronisé",latest_bytes,latest_name,mime,use_container_width=True)
        st.download_button("⬇️ Télécharger une copie Excel aplatie", export_flat_excel(st.session_state.df), "OPH65_suivi_travaux.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

elif page == "🏠 Tableau de bord":
    st.title("Tableau de bord")
    df = enrich(st.session_state.df)
    if df.empty:
        st.info("Importez d'abord le fichier de suivi.")
    else:
        counts = df["Statut global"].value_counts()
        cols = st.columns(5)
        labels = ["Logements suivis", "Terminés", "En cours", "En retard", "Non démarrés"]
        vals = [len(df), int(counts.get("Terminé",0)), int(counts.get("En cours",0)), int(counts.get("En retard",0)), int(counts.get("Non démarré",0))]
        for c, l, v in zip(cols, labels, vals): c.metric(l, v)
        left, right = st.columns(2)
        with left:
            fig = px.pie(df, names="Statut global", title="Répartition globale", color="Statut global", color_discrete_map=st.session_state.status_colors)
            st.plotly_chart(fig, use_container_width=True)
        with right:
            rows=[]
            for w in WORK_COLUMNS:
                for s,n in df[f"Statut {w}"].value_counts().items(): rows.append({"Lot":w,"Statut":s,"Nombre":n})
            chart=pd.DataFrame(rows)
            fig=px.bar(chart,x="Lot",y="Nombre",color="Statut",barmode="stack",title="Avancement par lot", color_discrete_map=st.session_state.status_colors, custom_data=["Lot","Statut"])
            event=st.plotly_chart(fig,use_container_width=True,on_select="rerun",selection_mode="points",key="dashboard_work_chart")
        try:
            points = event.selection.points
            if points:
                p = points[0]
                custom = p.get("customdata", [])
                if len(custom) >= 2: chart_detail_table(df, custom[0], custom[1])
        except Exception:
            st.caption("Cliquez sur une barre du diagramme « Avancement par lot » pour afficher les logements concernés.")

elif page == "📋 Suivi des logements":
    st.title("Suivi des logements")
    if st.session_state.df.empty: st.warning("Aucune donnée chargée.")
    else:
        df = enrich(st.session_state.df)
        show_cols=[c for c in REFERENCE_COLUMNS + [f"Statut {w}" for w in WORK_COLUMNS] + ["Statut global"] if c in df]
        show_excel_grid(df[show_cols], height=620, key="tracking_grid")

elif page == "🔎 Consultation des données":
    st.title("Consultation des données")
    if st.session_state.df.empty: st.warning("Aucune donnée chargée.")
    else:
        df = enrich(st.session_state.df)
        show_excel_grid(df, height=680, key="consult_grid")

elif page == "➕ Ajouter / Modifier":
    st.title("Ajouter / Modifier un logement")
    if st.session_state.df.empty:
        st.warning("Importez d'abord le fichier Excel afin de reprendre sa structure.")
    else:
        # La remise à zéro est exécutée AVANT la création des widgets, ce qui évite
        # l'erreur Streamlit liée à la modification d'un widget déjà instancié.
        if st.session_state.edit_form_reset_pending:
            st.session_state.edit_form_reset_pending = False
            form_prefixes = (
                "lot_no_", "gdc_", "ref_rsd_", "ref_bat_", "ref_ent_", "ref_porte_",
                "date_els_", "pal_", "order_", "start_", "planned_", "real_"
            )
            for k in list(st.session_state.keys()):
                if k == "edit_selected" or k.startswith(form_prefixes):
                    st.session_state.pop(k, None)
            st.session_state.edit_form_version += 1

        if st.session_state.save_dialog_pending:
            st.session_state.save_dialog_pending = False
            show_save_dialog()

        df = st.session_state.df
        lots = ["Nouveau logement"] + [str(x) for x in df["N° LOT"].dropna().tolist()] if "N° LOT" in df else ["Nouveau logement"]
        selected = st.selectbox("Logement à modifier", lots, key="edit_selected")
        is_new = selected == "Nouveau logement"
        existing = None if is_new else df[df["N° LOT"].astype(str) == selected].iloc[0]
        v = st.session_state.edit_form_version

        st.markdown("#### Identification du logement")
        c1,c2 = st.columns(2)
        lot_no = c1.text_input("N° de lot", value="" if existing is None else str(existing.get("N° LOT", "")), key=f"lot_no_{selected}_{v}")
        gdc = c2.text_input("GDC", value="" if existing is None else str(existing.get("GDC", "")), key=f"gdc_{selected}_{v}")

        lot_ref = lookup_lot_reference(lot_no)
        if lot_ref is None and existing is not None:
            # Compatibilité avec un ancien logement non encore présent dans la base.
            lot_ref = {
                "RSD": _clean(existing.get("RSD", "")),
                "Bât": _clean(existing.get("Bât", "")),
                "Ent": _clean(existing.get("Ent", "")),
                "Porte": _clean(existing.get("Porte", "")),
            }
        if lot_ref is None:
            lot_ref = {"RSD":"", "Bât":"", "Ent":"", "Porte":""}
            if normalize_lot_key(lot_no):
                st.warning("Ce N° LOT n'existe pas dans la base logements. Ajoutez-le dans Paramètres > Base logements pour renseigner automatiquement RSD, Bât, Entrée et Porte.")
        else:
            if normalize_lot_key(lot_no):
                st.success("Références logement récupérées automatiquement depuis la base OPH65.")

        rsd, bat, ent, porte = lot_ref.get("RSD", ""), lot_ref.get("Bât", ""), lot_ref.get("Ent", ""), lot_ref.get("Porte", "")
        c3,c4,c5,c6 = st.columns(4)
        c3.text_input("RSD", value=rsd, disabled=True, key=f"ref_rsd_{selected}_{v}_{normalize_lot_key(lot_no)}")
        c4.text_input("Bât", value=bat, disabled=True, key=f"ref_bat_{selected}_{v}_{normalize_lot_key(lot_no)}")
        c5.text_input("Entrée", value=ent, disabled=True, key=f"ref_ent_{selected}_{v}_{normalize_lot_key(lot_no)}")
        c6.text_input("Porte", value=porte, disabled=True, key=f"ref_porte_{selected}_{v}_{normalize_lot_key(lot_no)}")
        c7,c8 = st.columns(2)
        date_els = c7.date_input("Date saisie ELS", value=None if existing is None or pd.isna(existing.get("DATE SAISIE ELS", pd.NaT)) else pd.Timestamp(existing.get("DATE SAISIE ELS")).date(), format="DD/MM/YYYY", key=f"date_els_{selected}_{v}")
        pal = c8.date_input("Date de fin de travaux initiale / PAL", value=None if existing is None or pd.isna(existing.get("PAL", pd.NaT)) else pd.Timestamp(existing.get("PAL")).date(), format="DD/MM/YYYY", key=f"pal_{selected}_{v}")

        form_values = {}
        st.markdown("#### Travaux")
        tabs = st.tabs(list(WORK_COLUMNS.keys()))
        for tab, work in zip(tabs, WORK_COLUMNS):
            cols = WORK_COLUMNS[work]
            with tab:
                a,b,c,d = st.columns(4)
                order_val = "" if existing is None or pd.isna(existing.get(cols["order"], None)) else str(existing.get(cols["order"]))
                order = a.text_input("N° de bon de commande", value=order_val, key=f"order_{work}_{selected}_{v}")
                start_val = None if existing is None or pd.isna(existing.get(cols["start"], pd.NaT)) else pd.Timestamp(existing.get(cols["start"])).date()
                start_date = b.date_input("Date de saisie de la commande", value=start_val, key=f"start_{work}_{selected}_{v}", format="DD/MM/YYYY")
                planned = add_business_days(start_date, 5) if start_date else pd.NaT
                c.text_input("Prévision J+5 jours ouvrés (automatique)", value=format_date_fr(planned), disabled=True, key=f"planned_{work}_{selected}_{v}")
                real_val = None if existing is None or pd.isna(existing.get(cols["real"], pd.NaT)) else pd.Timestamp(existing.get(cols["real"])).date()
                real = d.date_input("Date réelle de fin des travaux", value=real_val, key=f"real_{work}_{selected}_{v}", format="DD/MM/YYYY")
                form_values[work] = {"order":order or None, "start":start_date, "planned":planned, "real":real}

        if st.button("💾 Enregistrer le logement", type="primary", key=f"save_housing_{v}"):
            if not lot_no.strip():
                st.error("Le N° de lot est obligatoire.")
            else:
                newrow = {c: None for c in df.columns}
                if existing is not None:
                    newrow.update(existing.to_dict())
                newrow.update({"N° LOT": lot_no.strip(), "GDC": gdc.strip() or None, "RSD": rsd or None, "Bât": bat or None, "Ent": ent or None, "Porte": porte or None, "DATE SAISIE ELS": date_els, "PAL": pal})
                for work, vals in form_values.items():
                    cols = WORK_COLUMNS[work]
                    newrow[cols["order"]] = vals["order"]
                    newrow[cols["start"]] = vals["start"]
                    newrow[cols["planned"]] = vals["planned"]
                    newrow[cols["real"]] = vals["real"]
                if is_new:
                    st.session_state.df = pd.concat([df, pd.DataFrame([newrow])], ignore_index=True)
                else:
                    idx = df[df["N° LOT"].astype(str) == selected].index[0]
                    for k,val in newrow.items():
                        if k in st.session_state.df.columns:
                            st.session_state.df.at[idx,k] = coerce_value_for_column(st.session_state.df, k, val)
                st.session_state.df = normalize_date_dtypes(st.session_state.df)
                st.session_state.df = apply_reference_database(st.session_state.df)
                st.session_state.df = calculate_planned_dates(st.session_state.df)
                if rsd and rsd not in st.session_state.residence_options:
                    st.session_state.residence_options = sorted(set(st.session_state.residence_options + [rsd]))
                    save_config(residences=st.session_state.residence_options)

                save_ok = True
                try:
                    if CLOUD_MODE and supabase_ready():
                        save_ok = persist_tracking_and_excel(st.session_state.df)
                    elif st.session_state.source_path and st.session_state.auto_save:
                        save_ok = persist_tracking_and_excel(st.session_state.df)
                except Exception as e:
                    save_ok = False
                    st.warning(f"Logement enregistré dans l'application, mais synchronisation persistante / Excel impossible : {e}")

                # Le formulaire est toujours remis à zéro après validation dans l'application.
                # Une erreur du classeur ne doit pas bloquer la saisie suivante.
                st.session_state.edit_form_reset_pending = True
                st.session_state.save_dialog_pending = True
                if save_ok:
                    st.session_state.save_dialog_message = "Enregistrement effectué. Les données persistantes et le classeur Excel synchronisé ont été mis à jour."
                elif st.session_state.source_path and st.session_state.auto_save:
                    st.session_state.save_dialog_message = "Enregistrement effectué dans l'application. Attention : le classeur Excel n'a pas pu être mis à jour ; rechargez le fichier original depuis Import / Export."
                else:
                    st.session_state.save_dialog_message = "Enregistrement effectué dans l'application."
                st.rerun()

elif page == "📦 Gestion des BC":
    render_bc_management()

elif page == "📊 Analyses":
    st.title("📊 Analyses")
    if st.session_state.df.empty: st.warning("Aucune donnée chargée.")
    else:
        df=enrich(st.session_state.df)
        choice=st.selectbox("Analyse", ["Retards par résidence","Avancement par lot","Liste des logements en retard","Répartition globale"])
        if choice=="Retards par résidence":
            x=df[df["Statut global"]=="En retard"]
            if x.empty: st.success("Aucun logement en retard.")
            else:
                chart=x.groupby("RSD",dropna=False).size().reset_index(name="Retards")
                st.plotly_chart(px.bar(chart,x="RSD",y="Retards",title="Retards par résidence"),use_container_width=True)
                show_excel_grid(x[[c for c in REFERENCE_COLUMNS+["Statut global"] if c in x]], height=380, key="analysis_delay")
        elif choice=="Avancement par lot":
            rows=[]
            for w in WORK_COLUMNS:
                for s,n in df[f"Statut {w}"].value_counts().items(): rows.append({"Lot":w,"Statut":s,"Nombre":n})
            st.plotly_chart(px.bar(pd.DataFrame(rows),x="Lot",y="Nombre",color="Statut",barmode="group",color_discrete_map=st.session_state.status_colors),use_container_width=True)
        elif choice=="Liste des logements en retard":
            show_excel_grid(df[df["Statut global"]=="En retard"], height=460, key="analysis_list")
        else:
            chart=df["Statut global"].value_counts().reset_index(); chart.columns=["Statut","Nombre"]
            st.plotly_chart(px.pie(chart,names="Statut",values="Nombre",color="Statut",color_discrete_map=st.session_state.status_colors),use_container_width=True)

        st.divider()
        st.markdown("### 📧 Synthèse d'avancement par email")
        synth = synthese_dataframe(st.session_state.df)
        show_excel_grid(synth, height=360, key="synth_grid")
        if HAS_REPORTLAB:
            st.download_button(
                "📄 Télécharger la synthèse PDF OPH65",
                data=export_synthese_pdf(st.session_state.df),
                file_name="Synthese_OPH65.pdf", mime="application/pdf",
                key="download_synth_pdf",
            )
        else:
            st.warning("ReportLab n'est pas installé : installez les dépendances pour générer le PDF de synthèse.")
        recipient = st.text_input("Destinataire(s)", value=st.session_state.email_recipients, key="analysis_recipient")
        default_subject = st.session_state.analysis_email_subject_template.replace("{date}", date.today().strftime("%d/%m/%Y"))
        subject = st.text_input("Objet", value=default_subject, key="analysis_subject")
        body = st.text_area("Message", value=st.session_state.analysis_email_body, height=180, key="analysis_body")
        cmemo, csend = st.columns([1, 1])
        if cmemo.button("💾 Mémoriser ce modèle de mail", use_container_width=True):
            st.session_state.email_recipients = recipient
            st.session_state.analysis_email_subject_template = subject.replace(date.today().strftime("%d/%m/%Y"), "{date}")
            st.session_state.analysis_email_body = body
            save_config(
                email_recipients=recipient,
                analysis_email_subject_template=st.session_state.analysis_email_subject_template,
                analysis_email_body=body,
            )
            st.success("Destinataires, objet et message mémorisés pour les prochains démarrages.")
        if csend.button("📤 Envoyer la synthèse", type="primary", use_container_width=True):
            if not (st.session_state.smtp_server and st.session_state.smtp_email and recipient) or (st.session_state.smtp_auth_required and not st.session_state.smtp_password):
                st.error("Configurez d'abord le SMTP dans Paramètres et renseignez un destinataire.")
            else:
                try:
                    send_email_with_attachment(st.session_state.smtp_server,st.session_state.smtp_port,st.session_state.smtp_email,st.session_state.smtp_password,recipient,subject,body,export_synthese_pdf(st.session_state.df),"Synthese_OPH65.pdf",st.session_state.smtp_use_ssl,st.session_state.smtp_auth_required,st.session_state.smtp_use_starttls)
                    st.success(f"Synthèse envoyée à {recipient}.")
                except smtplib.SMTPAuthenticationError as e:
                    st.error(smtp_auth_error_message(e))
                except Exception as e:
                    st.error(f"Échec de l'envoi : {e}")

elif page == "⚙️ Paramètres":
    st.title("⚙️ Paramètres")
    t1,t2,t3,t4,t5,t6,t7 = st.tabs(["🎨 Apparence & fond","🏠 Base logements","🏢 Résidences","📧 Destinataires par lot","👤 Utilisateurs / dates","✉️ Emails / SMTP","☁️ Base persistante"])

    with t1:
        st.markdown("### Apparence générale")
        c1, c2 = st.columns(2)
        p = c1.color_picker("Couleur principale", st.session_state.primary_color)
        sidebar = c2.color_picker("Couleur du menu", st.session_state.sidebar_color)
        cfont1, cfont2 = st.columns(2)
        font_options = ["Arial", "Verdana", "Tahoma", "Trebuchet MS", "Georgia", "Times New Roman", "Courier New"]
        current_font = st.session_state.font_family if st.session_state.font_family in font_options else "Arial"
        font_family = cfont1.selectbox("Police de caractères", font_options, index=font_options.index(current_font))
        font_color = cfont2.color_picker("Couleur du texte", st.session_state.font_color)

        st.markdown("#### Couleurs des statuts")
        status_cols = st.columns(4)
        new_status_colors = {}
        for col, status in zip(status_cols, STATUS_LIST):
            with col:
                new_status_colors[status] = st.color_picker(
                    status,
                    value=st.session_state.status_colors.get(status, DEFAULT_STATUS_COLORS[status]),
                    key=f"status_color_{status}",
                )

        st.divider()
        st.markdown("### Fond de l'application")
        bg_mode = st.radio("Type de fond", ["Couleur", "Image"], index=0 if st.session_state.background_mode == "Couleur" else 1, horizontal=True)
        bg_color = st.color_picker("Couleur de fond", value=st.session_state.background_color)
        st.markdown("### 🖼️ Image de fond")
        st.write(
            "Vous pouvez utiliser une image JPG, PNG ou WebP. Elle est conservée dans le dossier de l'application "
            "et s'adapte automatiquement à la taille et à la résolution de l'écran, sans déformation."
        )
        uploaded_bg = st.file_uploader("Choisir une image de fond", type=["jpg", "jpeg", "png", "webp"], key="background_uploader")
        overlay = st.slider(
            "Voile de lisibilité sur l'image",
            min_value=0.0,
            max_value=0.90,
            value=float(st.session_state.background_overlay),
            step=0.05,
            help="0 = image très visible ; 0,90 = fond très atténué pour faciliter la lecture.",
        )
        if st.session_state.background_image:
            current_bg = os.path.join(APP_DIR, os.path.basename(st.session_state.background_image))
            if os.path.exists(current_bg):
                st.caption("Image de fond actuellement enregistrée :")
                st.image(current_bg, width=420)

        ca, cb = st.columns(2)
        if ca.button("💾 Enregistrer l'apparence", type="primary", use_container_width=True):
            bg_name = st.session_state.background_image
            if uploaded_bg is not None:
                try:
                    bg_name = save_background_file(uploaded_bg)
                except Exception as e:
                    st.error(f"Impossible d'enregistrer l'image de fond : {e}")
                    bg_name = st.session_state.background_image
            st.session_state.primary_color = p
            st.session_state.sidebar_color = sidebar
            st.session_state.status_colors = new_status_colors
            st.session_state.background_image = bg_name
            st.session_state.background_overlay = overlay
            st.session_state.background_mode = bg_mode
            st.session_state.background_color = bg_color
            st.session_state.font_family = font_family
            st.session_state.font_color = font_color
            save_config(
                primary_color=p,
                sidebar_color=sidebar,
                status_colors=new_status_colors,
                background_image=bg_name,
                background_overlay=overlay,
                background_mode=bg_mode,
                background_color=bg_color,
                font_family=font_family,
                font_color=font_color,
            )
            st.success("Apparence enregistrée.")
            st.rerun()

        if cb.button("🗑️ Supprimer l'image de fond", use_container_width=True):
            delete_background_file(st.session_state.background_image)
            st.session_state.background_image = ""
            st.session_state.background_mode = "Couleur"
            save_config(background_image="", background_mode="Couleur")
            st.success("Image de fond supprimée.")
            st.rerun()

    with t2:
        st.markdown("### Base de référence des logements")
        st.write("Le **N° LOT est la clé unique**. Cette base alimente automatiquement les colonnes **RSD, Bât, Entrée et Porte** dans toute l'interface, les synthèses et les PDF.")
        reference_db = load_reference_database()
        m1, m2 = st.columns(2)
        m1.metric("Logements référencés", f"{len(reference_db):,}".replace(",", " "))
        m2.metric("Clé de référence", "N° LOT")

        search_ref = st.text_input("Rechercher dans la base", placeholder="N° lot, résidence, bâtiment, entrée ou porte", key="reference_db_search")
        ref_rows = list(reference_db.values())
        if search_ref.strip():
            q = search_ref.strip().lower()
            ref_rows = [r for r in ref_rows if any(q in str(r.get(c, "")).lower() for c in ["N° LOT","RSD","Bât","Ent","Porte"])]
        preview = pd.DataFrame(ref_rows[:500], columns=["N° LOT","RSD","Bât","Ent","Porte"]).rename(columns={"Ent":"Entrée"})
        st.caption(f"{len(ref_rows)} résultat(s). Affichage limité aux 500 premiers.")
        st.dataframe(preview, use_container_width=True, hide_index=True, height=330)

        st.markdown("#### Ajouter ou modifier une référence")
        edit_lot = st.text_input("N° LOT à ajouter / modifier", key="reference_db_edit_lot")
        edit_key = normalize_lot_key(edit_lot)
        current_ref = reference_db.get(edit_key, {}) if edit_key else {}
        r1,r2,r3,r4 = st.columns(4)
        ref_rsd = r1.text_input("RSD", value=current_ref.get("RSD", ""), key=f"db_edit_rsd_{edit_key or 'new'}")
        ref_bat = r2.text_input("Bât", value=current_ref.get("Bât", ""), key=f"db_edit_bat_{edit_key or 'new'}")
        ref_ent = r3.text_input("Entrée", value=current_ref.get("Ent", ""), key=f"db_edit_ent_{edit_key or 'new'}")
        ref_porte = r4.text_input("Porte", value=current_ref.get("Porte", ""), key=f"db_edit_porte_{edit_key or 'new'}")
        bsave, bdelete = st.columns(2)
        if bsave.button("💾 Enregistrer cette référence", type="primary", use_container_width=True):
            if not edit_key:
                st.error("Le N° LOT est obligatoire.")
            else:
                reference_db[edit_key] = {**current_ref, "N° LOT":edit_key, "RSD":ref_rsd.strip(), "Bât":ref_bat.strip(), "Ent":ref_ent.strip(), "Porte":ref_porte.strip()}
                save_reference_database(reference_db)
                if not st.session_state.df.empty:
                    st.session_state.df = apply_reference_database(st.session_state.df)
                if ref_rsd.strip() and ref_rsd.strip() not in st.session_state.residence_options:
                    st.session_state.residence_options = sorted(set(st.session_state.residence_options + [ref_rsd.strip()]))
                    save_config(residences=st.session_state.residence_options)
                st.success(f"Référence du lot {edit_key} enregistrée.")
                st.rerun()
        if bdelete.button("🗑️ Supprimer cette référence", use_container_width=True, disabled=not bool(edit_key and edit_key in reference_db)):
            reference_db.pop(edit_key, None)
            save_reference_database(reference_db)
            st.success(f"Référence du lot {edit_key} supprimée.")
            st.rerun()

        st.divider()
        st.markdown("#### Actualiser la base depuis un fichier OPH65")
        st.caption("Le fichier doit contenir l'onglet « List LGMT ». La fiche complète du logement sera importée : références, adresse, typologie, secteur, mise en service, chauffage et contrat PES.")
        ref_upload = st.file_uploader("Fichier Excel / XLSM de référence", type=["xlsx","xlsm","xls"], key="reference_db_upload")
        if st.button("📥 Remplacer la base par List LGMT", disabled=ref_upload is None, use_container_width=True):
            try:
                imported_db = reference_database_from_excel(ref_upload)
                save_reference_database(imported_db)
                if not st.session_state.df.empty:
                    st.session_state.df = apply_reference_database(st.session_state.df)
                st.success(f"Base actualisée : {len(imported_db)} logements importés.")
                st.rerun()
            except Exception as e:
                st.error(f"Impossible d'importer la base logements : {e}")

        export_cols = ["N° LOT","RSD","Bât","Ent","Porte","Résidence","Adresse","Suite adresse","Code postal","Ville","Typologie","Secteur","Date mise en service","Type chauffage","Contrat PES"]
        export_df = pd.DataFrame(reference_db.values(), columns=export_cols).rename(columns={"Ent":"Entrée"})
        export_csv = export_df.to_csv(index=False, sep=";", encoding="utf-8-sig").encode("utf-8-sig")
        st.download_button("⬇️ Exporter la base logements (CSV)", data=export_csv, file_name="Base_logements_OPH65.csv", mime="text/csv", use_container_width=True)

    with t3:
        current = "\n".join(st.session_state.residence_options)
        txt = st.text_area("Une résidence par ligne", value=current, height=300)
        if st.button("💾 Enregistrer les résidences"):
            st.session_state.residence_options = sorted(set(x.strip() for x in txt.splitlines() if x.strip()))
            save_config(residences=st.session_state.residence_options)
            st.success("Liste enregistrée et mémorisée pour les prochains démarrages.")

    with t4:
        emails = {}
        for w in WORK_COLUMNS:
            emails[w] = st.text_input(w, value=st.session_state.lot_emails.get(w,""), key=f"email_{w}")
        if st.button("💾 Enregistrer les destinataires"):
            st.session_state.lot_emails = emails
            save_config(lot_emails=emails)
            st.success("Destinataires enregistrés et mémorisés.")

    with t5:
        st.markdown("### Utilisateurs")
        f = st.selectbox("Format des dates", list(DATE_FORMAT_OPTIONS), index=list(DATE_FORMAT_OPTIONS).index(st.session_state.date_format))
        new_user = st.text_input("Nouvel utilisateur", placeholder="Nom et prénom")
        cadd, csave_user = st.columns(2)
        if cadd.button("➕ Créer l'utilisateur", use_container_width=True):
            name = new_user.strip()
            if not name:
                st.error("Saisissez un nom d'utilisateur.")
            elif name in st.session_state.users:
                st.warning("Cet utilisateur existe déjà.")
            else:
                st.session_state.users = sorted(set(st.session_state.users + [name]))
                save_config(users=st.session_state.users, username=st.session_state.username, date_format=f)
                st.success(f"Utilisateur « {name} » créé.")
                st.rerun()
        if csave_user.button("💾 Enregistrer le format des dates", use_container_width=True):
            st.session_state.date_format = f
            save_config(username=st.session_state.username, users=st.session_state.users, date_format=f)
            st.success("Paramètres enregistrés.")
        if len(st.session_state.users) > 1:
            remove_user = st.selectbox("Supprimer un utilisateur", [u for u in st.session_state.users if u != st.session_state.username])
            if st.button("🗑️ Supprimer l'utilisateur sélectionné"):
                st.session_state.users = [u for u in st.session_state.users if u != remove_user]
                save_config(users=st.session_state.users, username=st.session_state.username)
                st.success(f"Utilisateur « {remove_user} » supprimé.")
                st.rerun()

    with t6:
        if CLOUD_MODE:
            st.info("☁️ Mode Cloud : le mot de passe SMTP doit être enregistré dans Settings → Secrets de Streamlit (`[smtp] password = ...`). Il n'est jamais écrit dans les fichiers du dépôt.")
        st.warning("Pour Gmail / Google Workspace, utilisez un mot de passe d'application (si autorisé) et non le mot de passe habituel. En environnement OPH65, la solution recommandée est un relais SMTP fourni par le service informatique ; il peut être configuré avec ou sans authentification selon leurs consignes.")
        preset = st.selectbox("Configuration rapide", ["Personnalisée", "Gmail / Google Workspace (SSL 465)", "Gmail / Google Workspace (STARTTLS 587)", "Relais SMTP OPH65 / organisme"])
        if preset == "Gmail / Google Workspace (SSL 465)":
            st.session_state.smtp_server, st.session_state.smtp_port = "smtp.gmail.com", 465
            st.session_state.smtp_use_ssl, st.session_state.smtp_use_starttls, st.session_state.smtp_auth_required = True, False, True
        elif preset == "Gmail / Google Workspace (STARTTLS 587)":
            st.session_state.smtp_server, st.session_state.smtp_port = "smtp.gmail.com", 587
            st.session_state.smtp_use_ssl, st.session_state.smtp_use_starttls, st.session_state.smtp_auth_required = False, True, True
        a,b = st.columns(2)
        server = a.text_input("Serveur SMTP", value=st.session_state.smtp_server)
        sender = a.text_input("Adresse expéditeur", value=st.session_state.smtp_email)
        ssl_on = a.checkbox("SSL direct (généralement port 465)", value=st.session_state.smtp_use_ssl)
        auth_required = a.checkbox("Authentification par identifiant / mot de passe", value=st.session_state.smtp_auth_required, help="Décochez uniquement si votre service informatique vous fournit un relais SMTP OPH65 autorisé sans authentification.")
        use_starttls = a.checkbox("STARTTLS (généralement port 587)", value=st.session_state.smtp_use_starttls, disabled=ssl_on)
        port = b.number_input("Port", value=int(st.session_state.smtp_port), step=1)
        password = b.text_input(
            "Mot de passe d'application / mot de passe SMTP",
            value=st.session_state.smtp_password,
            type="password",
            help=("Sur Streamlit Cloud, configurez ce mot de passe dans Settings → Secrets. En local, il peut être mémorisé dans le coffre du système." if CLOUD_MODE else "Le mot de passe est mémorisé de manière sécurisée par le système, et non dans oph65_config.json."),
        )
        rec = b.text_input("Destinataire(s) par défaut", value=st.session_state.email_recipients)
        rts_email = b.text_input("Adresse mail Responsable Technique de Secteur", value=st.session_state.rts_email, help="Cette adresse reçoit automatiquement la validation de fin de prestation après scan du QR code.")

        st.markdown("#### Modèle du mail de synthèse")
        subject_template = st.text_input(
            "Objet par défaut",
            value=st.session_state.analysis_email_subject_template,
            help="Vous pouvez utiliser {date}, qui sera remplacé automatiquement par la date du jour.",
        )
        body_template = st.text_area("Message par défaut", value=st.session_state.analysis_email_body, height=190)

        st.markdown("#### Modèle du mail automatique de retard")
        st.caption("Variables disponibles : {lot}, {bc}, {rsd}, {bat}, {ent}, {porte}, {echeance}, {date}.")
        delay_subject_template = st.text_input("Objet du mail de retard", value=st.session_state.delay_email_subject_template)
        delay_body_template = st.text_area("Corps du mail de retard", value=st.session_state.delay_email_body_template, height=300)

        csave, ctest = st.columns(2)
        if csave.button("💾 Enregistrer tous les paramètres email", type="primary", use_container_width=True):
            st.session_state.smtp_server = server
            st.session_state.smtp_email = sender
            st.session_state.smtp_use_ssl = ssl_on
            st.session_state.smtp_auth_required = auth_required
            st.session_state.smtp_use_starttls = use_starttls
            st.session_state.smtp_port = int(port)
            st.session_state.smtp_password = password
            st.session_state.email_recipients = rec
            st.session_state.rts_email = rts_email.strip()
            st.session_state.analysis_email_subject_template = subject_template
            st.session_state.analysis_email_body = body_template
            st.session_state.delay_email_subject_template = delay_subject_template
            st.session_state.delay_email_body_template = delay_body_template
            save_config(
                smtp_server=server,
                smtp_email=sender,
                smtp_use_ssl=ssl_on,
                smtp_auth_required=auth_required,
                smtp_use_starttls=use_starttls,
                smtp_port=int(port),
                email_recipients=rec,
                rts_email=rts_email.strip(),
                analysis_email_subject_template=subject_template,
                analysis_email_body=body_template,
                delay_email_subject_template=delay_subject_template,
                delay_email_body_template=delay_body_template,
            )
            secure_ok = save_secure_smtp_password(sender, password)
            if CLOUD_MODE:
                if _secret("smtp.password", ""):
                    st.success("Paramètres email appliqués. Le secret SMTP est lu depuis Streamlit Secrets.")
                else:
                    st.warning("Paramètres appliqués pour cette session. Ajoutez `password` dans la section `[smtp]` des Secrets Streamlit pour conserver le secret de façon sûre.")
            elif password and not secure_ok:
                st.warning("Les paramètres ont été enregistrés, mais le système n'a pas permis de mémoriser le mot de passe. Il faudra le ressaisir au prochain démarrage.")
            else:
                st.success("Tous les paramètres email ont été enregistrés. Le mot de passe est mémorisé dans le coffre sécurisé du système.")

        if ctest.button("🧪 Tester l'authentification SMTP", use_container_width=True):
            if not (server and sender) or (auth_required and not password):
                st.error("Renseignez le serveur SMTP, l'adresse expéditeur et, si l'authentification est activée, le mot de passe SMTP.")
            else:
                try:
                    test_smtp_connection(server, port, sender, password, ssl_on, auth_required, use_starttls)
                    st.success("Connexion et authentification SMTP réussies.")
                except smtplib.SMTPAuthenticationError as e:
                    st.error(smtp_auth_error_message(e))
                except Exception as e:
                    st.error(f"Impossible de se connecter au serveur SMTP : {e}")



    with t7:
        st.markdown("### ☁️ Base de données persistante")
        if supabase_ready():
            st.success("Supabase est configuré et accessible depuis l'application.")
            try:
                cloud_df=db_load_tracking()
                ref_count=len(db_load_reference_database())
                wb_bytes,wb_name=db_load_excel_bytes()
                c1,c2,c3=st.columns(3)
                c1.metric("Logements suivis",len(cloud_df))
                c2.metric("Fiches logements",ref_count)
                c3.metric("Classeur Excel", "Disponible" if wb_bytes else "Non chargé")
                if wb_name: st.caption(f"Classeur synchronisé : {wb_name}")
            except Exception as e:
                st.warning(f"Connexion présente mais lecture impossible : {e}")
            if st.button("🔄 Recharger les données depuis Supabase",use_container_width=True):
                cloud_df=db_load_tracking()
                if not cloud_df.empty:
                    st.session_state.df=cloud_df
                    wb_bytes,wb_name=db_load_excel_bytes()
                    st.session_state.source_workbook_bytes=wb_bytes
                    st.session_state.source_workbook_name=wb_name
                    st.success("Données rechargées depuis Supabase.")
                    st.rerun()
        else:
            st.error("Supabase n'est pas encore configuré dans les Secrets Streamlit.")
            st.code('[supabase]\nurl = "https://VOTRE-PROJET.supabase.co"\nsecret_key = "VOTRE_CLE_SECRETE_SUPABASE"', language='toml')
            st.caption("Exécutez d'abord le fichier SUPABASE_SETUP.sql dans SQL Editor de Supabase, puis ajoutez ces deux valeurs dans Settings → Secrets de votre application Streamlit.")
